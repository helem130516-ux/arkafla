"""Gera ARKAFLA_Qualidade_do_Leite_2026.xlsx (openpyxl, fórmulas) a partir dos
relatórios APCBRH em dados/apcbrh e do registro de mastite em dados/mastite_clinica.json.

Uso: python3 scripts/gera_planilha.py [saida.xlsx]
Depois: recalcular com recalc.py (skill xlsx) e restaurar os XML dos gráficos
(scripts/restaura_graficos.py).
"""
import glob
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break

sys.path.insert(0, os.path.dirname(__file__))
from processa_controles import classifica, le_relatorio  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else "saida/ARKAFLA_Qualidade_do_Leite_2026.xlsx"
F = "Arial"
AZUL = "1F5F8B"
COR = dict(S="2E8B57", R="9AA39D", PB="E8C547", PE="E8833A", H="FFFFFF", N="1B3A6B", C="C0392B")
NOME = dict(S="Sadia", N="Nova infecção", C="Crônica", R="Recuperada",
            PB="Pós-parto baixa", PE="Pós-parto elevada", H="Sem histórico")

thin = Side(style="thin", color="C9D1CB")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
HFILL = PatternFill("solid", fgColor="E6ECE8")
INFILL = PatternFill("solid", fgColor="FFF6D6")
SECFILL = PatternFill("solid", fgColor=AZUL)


def font(**k):
    return Font(name=F, **{"size": 10, **k})


def hdr(ws, row, cols, start=1, fill=HFILL, wrap=True):
    for i, t in enumerate(cols):
        c = ws.cell(row, start + i, t)
        c.font = font(bold=True, size=9)
        c.fill = fill
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14, color=AZUL)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, size=9, color="5E6B63")


def sec(ws, row, text, ncol):
    c = ws.cell(row, 1, text)
    c.font = font(bold=True, color="FFFFFF")
    for j in range(1, ncol + 1):
        ws.cell(row, j).fill = SECFILL


def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


# ---------------------------------------------------------------- dados
ctrls = [le_relatorio(p) for p in sorted(glob.glob("dados/apcbrh/R22_*.xlsx"))]
recs = classifica(ctrls)
NC = len(ctrls)
mast = json.load(open("dados/mastite_clinica.json"))
mast_rows = mast["tratamentos"]
NR = len(recs)
NM = len(mast_rows)
DL = NR + 1            # última linha de Dados
ML = NM + 1            # última linha de Mastite
vacas = sorted({int(x["vaca"]) for x in recs} | {int(m["vaca"]) for m in mast_rows})
NV = len(vacas)
LOTES = sorted({int(x["lote"]) for x in recs})
CL = [L(3 + i) for i in range(NC)]  # colunas por controle no Painel: C..K

wb = Workbook()
wb._named_styles["Normal"].font = Font(name=F, size=10)
wb.remove(wb.active)
order = ["Leia-me", "Painel", "Impressão", "Gráficos", "Lotes", "Ranking", "Histórico",
         "Ficha da vaca", "Mastite clínica", "Cultura", "Critérios", "Qualidade dos dados",
         "Parâmetros", "Controles", "Dados", "Calc", "Dados gráficos"]
S = {n: wb.create_sheet(n) for n in order}


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def R(sheet, col, a=2, b=None):
    """Intervalo absoluto numa aba."""
    b = b or {"Dados": DL, "Mastite clínica": ML}[sheet]
    return f"'{sheet}'!${col}${a}:${col}${b}"


# ---------------------------------------------------------------- Parâmetros
ws = S["Parâmetros"]
title(ws, "Parâmetros da análise", "Células amarelas podem ser alteradas; todas as abas recalculam.")
params = [
    ("LIM", "Limite individual de CCS (mil céls/mL)", 200),
    ("FX2", "Faixa 2 começa em", 200),
    ("FX3", "Faixa 3 começa em", 500),
    ("FX4", "Faixa 4 começa em", 1000),
    ("QUEDA", "Queda de produção (fração do controle anterior)", 0.30),
    ("DELMIN", "DEL mínimo para contar queda de produção", 60),
    ("CISEL", "Controle analisado (Nº do controle, 1 a %d)" % NC, NC),
    ("TOPN", "Top N para a simulação de descarte (1 a 50)", 20),
    ("DESVIO", "Lote desviado (tratamento)", 10),
    ("RETDIAS", "Retratamento: dias desde o início anterior", 21),
    ("TETO", "CCS considerada teto de leitura", 9998),
    ("INTMIN", "Intervalo mínimo (dias) para taxa por 30 dias", 28),
]
hdr(ws, 4, ["Parâmetro", "Valor", "Nome"])
for i, (n, t, v) in enumerate(params):
    r = 5 + i
    ws.cell(r, 1, t).font = font()
    c = ws.cell(r, 2, v)
    c.fill = INFILL
    c.font = font(bold=True, color="0000FF")
    c.border = BOX
    ws.cell(r, 3, n).font = font(size=9, color="5E6B63")
    name(n, f"'Parâmetros'!$B${r}")
ws.cell(5 + 4, 2).number_format = "0%"
r = 5 + len(params) + 1
ws.cell(r, 1, "Controle analisado").font = font(bold=True)
ws.cell(r, 2, "=TEXT(INDEX(Controles!$B$2:$B$%d,CISEL),\"DD/MM/YYYY\")" % (NC + 1)).font = font(bold=True)
widths(ws, {"A": 52, "B": 14, "C": 10})

# ---------------------------------------------------------------- Controles
ws = S["Controles"]
cols = ["Nº controle", "Data do controle", "REB", "Nº animais (APCBRH)", "MÉDIA CCS TANQUE (APCBRH)",
        "Produção total (APCBRH)", "Média produção (APCBRH)", "Dias desde o anterior",
        "Vacas carregadas", "Σ produção (kg)", "Σ produção × CCS", "CCS ponderada calculada",
        "Diferença calc − APCBRH", "Confere?", "Rótulo", "Arquivo", "Data de emissão"]
hdr(ws, 1, cols)
for i, c in enumerate(ctrls):
    r = 2 + i
    vals = [c["ci"], c["data"], c["reb"], c["n"], c["ccs_tq"], c["prod_total"], c["media_prod"]]
    for j, v in enumerate(vals):
        ws.cell(r, 1 + j, v)
    ws.cell(r, 2).number_format = "DD/MM/YYYY"
    ws.cell(r, 8, f"=B{r}-B{r-1}" if i else "")
    ws.cell(r, 9, f"=COUNTIF({R('Dados','B')},A{r})")
    ws.cell(r, 10, f"=SUMIF({R('Dados','B')},A{r},{R('Dados','H')})")
    ws.cell(r, 11, f"=SUMIF({R('Dados','B')},A{r},{R('Dados','T')})")
    ws.cell(r, 12, f"=K{r}/J{r}")
    ws.cell(r, 13, f"=L{r}-E{r}")
    ws.cell(r, 14, f'=IF(AND(ABS(M{r})<0.05,I{r}=D{r}),"OK","VERIFICAR")')
    ws.cell(r, 15, f'=TEXT(B{r},"DD/MM")')
    ws.cell(r, 16, c["arquivo"])
    ws.cell(r, 17, c["emissao"])
    for j in (5, 7, 12):
        ws.cell(r, j).number_format = "0.0"
    ws.cell(r, 6).number_format = "#,##0.0"
    ws.cell(r, 10).number_format = "#,##0.0"
    ws.cell(r, 11).number_format = "#,##0"
    ws.cell(r, 13).number_format = "0.000"
widths(ws, {L(i): w for i, w in enumerate([10, 12, 8, 11, 13, 13, 11, 10, 10, 12, 14, 12, 12, 10, 8, 22, 12], 1)})
ws.freeze_panes = "B2"
CTR = f"Controles!$B$2:$B${NC+1}"

# ---------------------------------------------------------------- Dados
ws = S["Dados"]
cols = ["Vaca", "Nº controle", "Data controle", "Lote", "Parto", "DEL", "Idade", "Produção (kg)",
        "CCS (mil/mL)", "Elevada (≥ limite)", "Mesma vaca que a linha acima", "Nova lactação",
        "Controle consecutivo", "CCS anterior", "Classe", "Já teve CCS elevada nesta lactação",
        "Controles seguidos elevada", "Recorrente", "Queda de produção", "Produção × CCS",
        "Lote desviado", "Faixa", "Parto alterado", "Chave ranking CCS", "Chave ranking contribuição",
        "ln(CCS)", "Parto em branco", "CCS no teto", "Anterior elevada", "Chave vaca-controle",
        "Presente no controle anterior", "CCS/TQ (APCBRH)"]
hdr(ws, 1, cols)
for i, x in enumerate(recs):
    r = 2 + i
    ws.cell(r, 1, int(x["vaca"]))
    ws.cell(r, 2, x["ci"])
    ws.cell(r, 3, f"=INDEX({CTR},B{r})")
    ws.cell(r, 4, int(x["lote"]))
    if x["parto"]:
        ws.cell(r, 5, x["parto"])
    ws.cell(r, 6, x["del_"])
    ws.cell(r, 7, x["idade"])
    ws.cell(r, 8, x["prod"])
    ws.cell(r, 9, x["ccs"])
    ws.cell(r, 10, f"=IF(I{r}>=LIM,1,0)")
    ws.cell(r, 11, f"=IF(A{r}=A{r-1},1,0)")
    ws.cell(r, 12, f'=IF(K{r}=1,IF(E{r}<>"",IF(E{r}>C{r-1},1,0),0),0)')
    ws.cell(r, 13, f"=IF(K{r}=1,IF(B{r}=B{r-1}+1,1,0),0)")
    ws.cell(r, 14, f'=IF(K{r}=1,I{r-1},"")')
    ws.cell(r, 15, f'=IF(K{r}=0,"H",IF(L{r}=1,IF(J{r}=1,"PE","PB"),IF(M{r}=0,"H",'
                   f'IF(J{r-1}=0,IF(J{r}=0,"S","N"),IF(J{r}=1,"C","R")))))')
    ws.cell(r, 16, f"=IF(K{r}=1,IF(L{r}=0,MAX(P{r-1},J{r-1}),0),0)")
    ws.cell(r, 17, f"=IF(J{r}=0,0,IF(M{r}=1,IF(L{r}=0,Q{r-1}+1,1),1))")
    ws.cell(r, 18, f'=IF(O{r}="N",IF(P{r}=1,1,0),0)')
    ws.cell(r, 19, f"=IF(J{r}=1,IF(K{r}=1,IF(L{r}=0,IF(H{r-1}>0,IF(H{r}<=(1-QUEDA)*H{r-1},"
                   f"IF(F{r}>=DELMIN,1,0),0),0),0),0),0)")
    ws.cell(r, 20, f"=H{r}*I{r}")
    ws.cell(r, 21, f"=IF(D{r}=DESVIO,1,0)")
    ws.cell(r, 22, f'=IF(I{r}<FX2,"1: < "&FX2,IF(I{r}<FX3,"2: "&FX2&"–"&(FX3-1),'
                   f'IF(I{r}<FX4,"3: "&FX3&"–"&(FX4-1),"4: ≥ "&FX4)))')
    ws.cell(r, 23, f'=IF(K{r}=1,IF(L{r}=0,IF(E{r}<>"",IF(E{r-1}<>"",IF(E{r}<>E{r-1},1,0),0),0),0),0)')
    ws.cell(r, 24, f'=IF(B{r}=CISEL,I{r}+ROW()/10000000,"")')
    ws.cell(r, 25, f'=IF(B{r}=CISEL,IF(U{r}=0,T{r}+ROW()/10000000,""),"")')
    ws.cell(r, 26, f"=LN(MAX(I{r},1))")
    ws.cell(r, 27, f'=IF(OR(E{r}="",F{r}=0),1,0)')
    ws.cell(r, 28, f"=IF(I{r}>=TETO,1,0)")
    ws.cell(r, 29, f'=IF(K{r}=1,J{r-1},"")')
    ws.cell(r, 30, f"=A{r}*100+B{r}")
    ws.cell(r, 31, f"=M{r}")
    ws.cell(r, 32, x["ccs_tq"])
    ws.cell(r, 3).number_format = "DD/MM/YYYY"
    ws.cell(r, 5).number_format = "DD/MM/YYYY"
widths(ws, {L(i): 9 for i in range(1, 33)})
ws.column_dimensions["V"].width = 12
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:AF{DL}"
D = lambda col: R("Dados", col)  # noqa: E731

# ---------------------------------------------------------------- Mastite clínica
ws = S["Mastite clínica"]
cols = ["Vaca", "Início tratamento", "Intramamário (produto)", "Quartos", "Data final",
        "Carência (dias)", "Retorno do leite", "Nº controle antes", "CCS antes",
        "Nº controle depois", "CCS depois", "Retratamento (≤ dias)", "Lote do caso",
        "Controle do período", "Sem produto", "Sem quarto", "CCS antes ≥ limite", "Ordem na ficha",
        "Vaca nos controles", "Observação do registro"]
hdr(ws, 1, cols)
for i, m in enumerate(mast_rows):
    r = 2 + i
    ws.cell(r, 1, int(m["vaca"]))
    ws.cell(r, 2, date.fromisoformat(m["ini"]))
    ws.cell(r, 3, m["produto"])
    ws.cell(r, 4, m["quartos"] or None)
    if m["fim"]:
        ws.cell(r, 5, date.fromisoformat(m["fim"]))
    ws.cell(r, 6, m["carencia"])
    ws.cell(r, 7, date.fromisoformat(m["retorno"]) if m.get("retorno") else f'=IF(E{r}="","",E{r}+F{r})')
    ws.cell(r, 8, f"=_xlfn.MAXIFS({D('B')},{D('A')},A{r},{D('C')},\"<=\"&B{r})")
    ws.cell(r, 9, f'=IF(H{r}=0,"",SUMIFS({D("I")},{D("A")},A{r},{D("B")},H{r}))')
    ws.cell(r, 10, f'=IF(G{r}="",0,_xlfn.MINIFS({D("B")},{D("A")},A{r},{D("C")},">"&G{r}))')
    ws.cell(r, 11, f'=IF(J{r}=0,"",SUMIFS({D("I")},{D("A")},A{r},{D("B")},J{r}))')
    ws.cell(r, 12, f"=IF(COUNTIFS($A$2:$A${ML},A{r},$B$2:$B${ML},\">=\"&(B{r}-RETDIAS),$B$2:$B${ML},\"<\"&B{r})"
                   f"+COUNTIFS($A$1:A{r-1},A{r},$B$1:B{r-1},B{r})>0,1,0)" if r > 2 else
                   f"=IF(COUNTIFS($A$2:$A${ML},A{r},$B$2:$B${ML},\">=\"&(B{r}-RETDIAS),$B$2:$B${ML},\"<\"&B{r})>0,1,0)")
    ws.cell(r, 13, f'=IF(H{r}=0,"",SUMIFS({D("D")},{D("A")},A{r},{D("B")},H{r}))')
    ws.cell(r, 14, f'=IF(B{r}>MAX({CTR}),"",COUNTIF({CTR},"<"&B{r})+1)')
    ws.cell(r, 15, f'=IF(C{r}="(sem produto)",1,0)')
    ws.cell(r, 16, f'=IF(D{r}="",1,0)')
    ws.cell(r, 17, f'=IF(I{r}="","",IF(I{r}>=LIM,1,0))')
    ws.cell(r, 18, f"=IF(A{r}='Ficha da vaca'!$C$3,COUNTIF($A$2:A{r},A{r}),\"\")")
    ws.cell(r, 19, f'=IF(COUNTIF({D("A")},A{r})>0,"sim","NÃO")')
    ws.cell(r, 20, m.get("obs") or None)
    for j in (2, 5, 7):
        ws.cell(r, j).number_format = "DD/MM/YYYY"
widths(ws, {L(i): w for i, w in enumerate([8, 12, 18, 10, 12, 9, 12, 9, 9, 9, 9, 11, 8, 9, 8, 8, 9, 8, 9, 34], 1)})
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:T{ML}"
M = lambda col: R("Mastite clínica", col)  # noqa: E731

# ---------------------------------------------------------------- Calc (agregados por controle)
ws = S["Calc"]
title(ws, "Cálculos auxiliares por controle", "Base para o Painel, Gráficos e Impressão. Não editar.")
ccols = [
    ("ci", None), ("Data", None), ("Vacas", "=COUNTIF({B},{ci})"),
    ("Σ prod", "=SUMIF({B},{ci},{H})"), ("Σ prod×CCS", "=SUMIF({B},{ci},{T})"),
    ("Σ CCS", "=SUMIF({B},{ci},{I})"), ("Σ ln CCS", "=SUMIF({B},{ci},{Z})"),
    ("Vacas sem lote desv.", "=COUNTIFS({B},{ci},{U},0)"),
    ("Σ prod sem desv.", "=SUMIFS({H},{B},{ci},{U},0)"),
    ("Σ prod×CCS sem desv.", "=SUMIFS({T},{B},{ci},{U},0)"),
    ("Σ CCS sem desv.", "=SUMIFS({I},{B},{ci},{U},0)"),
    ("Σ ln sem desv.", "=SUMIFS({Z},{B},{ci},{U},0)"),
    ("S", '=COUNTIFS({B},{ci},{O},"S")'), ("N", '=COUNTIFS({B},{ci},{O},"N")'),
    ("C", '=COUNTIFS({B},{ci},{O},"C")'), ("R", '=COUNTIFS({B},{ci},{O},"R")'),
    ("PB", '=COUNTIFS({B},{ci},{O},"PB")'), ("PE", '=COUNTIFS({B},{ci},{O},"PE")'),
    ("H", '=COUNTIFS({B},{ci},{O},"H")'),
    ("Recorrentes", "=COUNTIFS({B},{ci},{R},1)"),
    ("Seguidos ≥ 2", '=COUNTIFS({B},{ci},{Q},">=2")'), ("Seguidos ≥ 3", '=COUNTIFS({B},{ci},{Q},">=3")'),
    ("Queda prod.", "=COUNTIFS({B},{ci},{S},1)"),
    ("≥ lim", '=COUNTIFS({B},{ci},{I},">="&LIM)'),
    ("≥ FX3", '=COUNTIFS({B},{ci},{I},">="&FX3)'), ("≥ FX4", '=COUNTIFS({B},{ci},{I},">="&FX4)'),
    ("Faixa 1", '=COUNTIFS({B},{ci},{I},"<"&FX2)'),
    ("Faixa 2", '=COUNTIFS({B},{ci},{I},">="&FX2,{I},"<"&FX3)'),
    ("Faixa 3", '=COUNTIFS({B},{ci},{I},">="&FX3,{I},"<"&FX4)'),
    ("Faixa 4", '=COUNTIFS({B},{ci},{I},">="&FX4)'),
    ("PB ant. elevada", '=COUNTIFS({B},{ci},{O},"PB",{AC},1)'),
    ("PE ant. elevada", '=COUNTIFS({B},{ci},{O},"PE",{AC},1)'),
    ("PB ant. baixa", '=COUNTIFS({B},{ci},{O},"PB",{AC},0)'),
    ("PE ant. baixa", '=COUNTIFS({B},{ci},{O},"PE",{AC},0)'),
    ("Partos em branco", "=COUNTIFS({B},{ci},{AA},1)"), ("CCS no teto", "=COUNTIFS({B},{ci},{AB},1)"),
    ("Partos alterados", "=COUNTIFS({B},{ci},{W},1)"),
    ("Presentes no anterior", "=COUNTIFS({B},{ci},{AE},1)"),
    ("Vacas no lote desviado", "=COUNTIFS({B},{ci},{U},1)"),
    ("Tratamentos no período", "=COUNTIF({MN},{ci})"),
    ("Retratamentos no período", "=COUNTIFS({MN},{ci},{ML},1)"),
    ("Casos com CCS antes", '=COUNTIFS({MN},{ci},{MQ},"<>")'),
    ("Casos com CCS antes ≥ lim", "=COUNTIFS({MN},{ci},{MQ},1)"),
    ("Casos sem produto", "=COUNTIFS({MN},{ci},{MO},1)"),
    ("Σ prod×CCS/TQ", None),
]
hdr(ws, 3, [c[0] for c in ccols[:-1]])
refs = {k: D(k) for k in "BHIOQRSTUZ"}
refs.update(W=D("W"), AA=D("AA"), AB=D("AB"), AC=D("AC"), AE=D("AE"),
            MN=M("N"), ML=M("L"), MQ=M("Q"), MO=M("O"))
CALC = {}  # nome -> letra da coluna em Calc
for i in range(NC):
    r = 4 + i
    for j, (nm, f) in enumerate(ccols[:-1]):
        col = L(1 + j)
        CALC[nm] = col
        if nm == "ci":
            ws.cell(r, 1, i + 1)
        elif nm == "Data":
            ws.cell(r, 2, f"=INDEX({CTR},A{r})").number_format = "DD/MM"
        else:
            ws.cell(r, 1 + j, f.format(ci=f"$A{r}", **refs))
widths(ws, {L(i): 10 for i in range(1, len(ccols))})
ws.freeze_panes = "C4"


def cref(nm, i):
    """Referência à Calc para o controle de índice i (0-based)."""
    return f"Calc!${CALC[nm]}${4 + i}"


def csel(nm, off=0):
    """Referência à Calc para o controle selecionado (CISEL + off)."""
    return f"INDEX(Calc!${CALC[nm]}$4:${CALC[nm]}${3 + NC},CISEL{'+%d' % off if off > 0 else ('-%d' % -off if off < 0 else '')})"


# ranking de contribuição (50 linhas) para a simulação de descarte, em Calc
RK0 = 4 + NC + 3
ws.cell(RK0 - 1, 1, "Top 50 contribuições (controle selecionado, sem lote desviado)").font = font(bold=True)
hdr(ws, RK0, ["k", "Chave", "Linha em Dados", "Produção", "Prod×CCS"])
for k in range(1, 51):
    r = RK0 + k
    ws.cell(r, 1, k)
    ws.cell(r, 2, f"=IFERROR(LARGE({D('Y')},A{r}),\"\")")
    ws.cell(r, 3, f'=IF(B{r}="","",MATCH(B{r},{D("Y")},0)+1)')
    ws.cell(r, 4, f'=IF(C{r}="",0,INDEX(Dados!$H:$H,C{r}))')
    ws.cell(r, 5, f'=IF(C{r}="",0,INDEX(Dados!$T:$T,C{r}))')
RK1, RK2 = RK0 + 1, RK0 + 50

# ---------------------------------------------------------------- Painel
ws = S["Painel"]
title(ws, "Painel de indicadores – Fazenda ARKAFLA (Castro/PR)",
      "Controle leiteiro APCBRH, Relatório 2.2 · CCS em mil células/mL · estimativas do tanque = Σ(produção × CCS) ÷ Σ produção no dia do controle")
hdr(ws, 4, ["Indicador", "Unid."] + [f"=Controles!O{2+i}" for i in range(NC)])
ws.cell(3, 3, "Nº controle →").font = font(size=8, color="5E6B63")
for i in range(NC):
    ws.cell(3 + 0, 3 + i).value = None
row = 5
PROW = {}


def prow(key, label, unit, fn, fmt="0.0", first_blank=False):
    """fn(i) -> fórmula para o controle i (0-based)."""
    global row
    ws.cell(row, 1, label).font = font()
    ws.cell(row, 2, unit).font = font(size=8, color="5E6B63")
    for i in range(NC):
        c = ws.cell(row, 3 + i, '=""' if (first_blank and i == 0) else fn(i))
        c.number_format = fmt
        c.border = BOX
        c.font = font()
    PROW[key] = row
    row += 1


def c(nm, i):
    return cref(nm, i)


def pdiv(a, b):
    return f'=IF({b}=0,"",{a}/{b})'


BASE = lambda i: f"({c('S',i)}+{c('N',i)}+{c('C',i)}+{c('R',i)})"  # noqa: E731
sec(ws, row, "10 indicadores do projeto", 2 + NC); row += 1
prow("i1", "1. CCS estimada do tanque sem o lote desviado (estimativa)", "mil/mL",
     lambda i: f"={c('Σ prod×CCS sem desv.',i)}/{c('Σ prod sem desv.',i)}", "0")
prow("i2", "2. % sadias (baixa → baixa)", "%", lambda i: pdiv(c('S', i), BASE(i)), "0.0%", True)
prow("i3", "3. % novas infecções (baixa → elevada)", "%", lambda i: pdiv(c('N', i), BASE(i)), "0.0%", True)
prow("i4", "4. % crônicas (elevada → elevada)", "%", lambda i: pdiv(c('C', i), BASE(i)), "0.0%", True)
prow("i5", "5. % recuperadas (elevada → baixa)", "%", lambda i: pdiv(c('R', i), BASE(i)), "0.0%", True)
prow("i6", "6. % vacas com CCS ≥ limite", "%", lambda i: f"={c('≥ lim',i)}/{c('Vacas',i)}", "0.0%")
prow("i7", "7. % vacas com CCS ≥ 500", "%", lambda i: f"={c('≥ FX3',i)}/{c('Vacas',i)}", "0.0%")
prow("i8", "8. % vacas com CCS ≥ 1.000", "%", lambda i: f"={c('≥ FX4',i)}/{c('Vacas',i)}", "0.0%")
prow("i9", "9. Produção média", "kg/vaca/dia", lambda i: f"={c('Σ prod',i)}/{c('Vacas',i)}", "0.0")
prow("i10", "10. Vacas com CCS elevada em ≥ 2 controles seguidos", "vacas",
     lambda i: f"={c('Seguidos ≥ 2',i)}", "0", True)
row += 1
sec(ws, row, "A. Rebanho", 2 + NC); row += 1
prow("n", "Vacas avaliadas (em lactação)", "vacas", lambda i: f"={c('Vacas',i)}", "#,##0")
prow("dias", "Dias desde o controle anterior", "dias", lambda i: f"=Controles!H{2+i}", "0", True)
prow("reb", "REB", "", lambda i: f"=Controles!C{2+i}", "0")
prow("apc", '"MÉDIA CCS TANQUE" APCBRH, com o lote desviado (estimativa)', "mil/mL",
     lambda i: f"=Controles!E{2+i}", "0")
prow("pond", "CCS ponderada recalculada, todas as vacas (conferência)", "mil/mL",
     lambda i: f"={c('Σ prod×CCS',i)}/{c('Σ prod',i)}", "0.0")
prow("arit", "CCS média aritmética sem o lote desviado", "mil/mL",
     lambda i: f"={c('Σ CCS sem desv.',i)}/{c('Vacas sem lote desv.',i)}", "0")
prow("geo", "CCS média geométrica sem o lote desviado", "mil/mL",
     lambda i: f"=EXP({c('Σ ln sem desv.',i)}/{c('Vacas sem lote desv.',i)})", "0")
prow("n10", "Vacas no lote desviado", "vacas", lambda i: f"={c('Vacas no lote desviado',i)}", "0")
prow("nel", "Vacas com CCS ≥ limite", "vacas", lambda i: f"={c('≥ lim',i)}", "0")
prow("f1", "Faixa < 200", "vacas", lambda i: f"={c('Faixa 1',i)}", "0")
prow("f2", "Faixa 200–499", "vacas", lambda i: f"={c('Faixa 2',i)}", "0")
prow("f3", "Faixa 500–999", "vacas", lambda i: f"={c('Faixa 3',i)}", "0")
prow("f4", "Faixa ≥ 1.000", "vacas", lambda i: f"={c('Faixa 4',i)}", "0")
prow("ent", "Vacas que entraram (não estavam no controle anterior)", "vacas",
     lambda i: f"={c('Vacas',i)}-{c('Presentes no anterior',i)}", "0", True)
prow("sai", "Vacas que saíram (estavam no anterior e não neste)", "vacas",
     lambda i: f"={c('Vacas',i-1)}-{c('Presentes no anterior',i)}" if i else '=""', "0", True)
row += 1
sec(ws, row, "B. Saúde da glândula mamária (base: vacas com CCS neste controle e no anterior)", 2 + NC); row += 1
prow("base", "Vacas na base da classificação", "vacas", lambda i: f"={BASE(i)}", "#,##0", True)
for k in "SNCR":
    prow("n" + k, f"{NOME[k]} (nº)", "vacas", lambda i, k=k: f"={c(k,i)}", "0", True)
prow("tni", "Taxa de novas infecções = NI ÷ (Sadias + NI)", "%",
     lambda i: pdiv(c('N', i), f"({c('S',i)}+{c('N',i)})"), "0.0%", True)
prow("tni30", "Taxa de novas infecções por 30 dias (estimativa; intervalo ≥ 28 d)", "%",
     lambda i: f'=IF(Controles!H{2+i}<INTMIN,"",{c("N",i)}/({c("S",i)}+{c("N",i)})*30/Controles!H{2+i})', "0.0%", True)
prow("perm", "Permanência = Crônicas ÷ (Crônicas + Recuperadas)", "%",
     lambda i: pdiv(c('C', i), f"({c('C',i)}+{c('R',i)})"), "0.0%", True)
prow("rec", "Recorrentes (NI em vaca que já teve CCS elevada na lactação)", "vacas",
     lambda i: f"={c('Recorrentes',i)}", "0", True)
prow("recp", "Recorrência = recorrentes ÷ NI", "%", lambda i: pdiv(c('Recorrentes', i), c('N', i)), "0.0%", True)
prow("seq3", "CCS elevada em ≥ 3 controles seguidos", "vacas", lambda i: f"={c('Seguidos ≥ 3',i)}", "0", True)
prow("queda", "CCS elevada + queda de produção ≥ 30% (DEL ≥ 60)", "vacas", lambda i: f"={c('Queda prod.',i)}", "0", True)
prow("PB", "Pós-parto baixa (fora da base)", "vacas", lambda i: f"={c('PB',i)}", "0", True)
prow("PE", "Pós-parto elevada (fora da base)", "vacas", lambda i: f"={c('PE',i)}", "0", True)
prow("H", "Sem histórico (fora da base)", "vacas", lambda i: f"={c('H',i)}", "0")
prow("cura", "Cura no período seco (elevada antes → baixa pós-parto)", "%",
     lambda i: pdiv(c('PB ant. elevada', i), f"({c('PB ant. elevada',i)}+{c('PE ant. elevada',i)})"), "0.0%", True)
prow("curan", "   vacas elevadas antes da secagem (n)", "vacas",
     lambda i: f"={c('PB ant. elevada',i)}+{c('PE ant. elevada',i)}", "0", True)
prow("nis", "Nova infecção no período seco (baixa antes → elevada pós-parto)", "%",
     lambda i: pdiv(c('PE ant. baixa', i), f"({c('PB ant. baixa',i)}+{c('PE ant. baixa',i)})"), "0.0%", True)
prow("nisn", "   vacas baixas antes da secagem (n)", "vacas",
     lambda i: f"={c('PB ant. baixa',i)}+{c('PE ant. baixa',i)}", "0", True)
row += 1
sec(ws, row, "C. Mastite clínica (registro da fazenda; tratamentos iniciados desde o controle anterior)", 2 + NC); row += 1
prow("trat", "Tratamentos no período", "trat.", lambda i: f"={c('Tratamentos no período',i)}", "0", True)
prow("retr", "   retratamentos (mesma vaca em ≤ 21 dias)", "trat.",
     lambda i: f"={c('Retratamentos no período',i)}", "0", True)
prow("t100", "Tratamentos por 100 vacas a cada 30 dias", "trat.",
     lambda i: f"={c('Tratamentos no período',i)}/{c('Vacas',i)}*100*30/Controles!H{2+i}", "0.0", True)
prow("tant", "Casos em vacas já com CCS ≥ limite no controle anterior", "%",
     lambda i: pdiv(c('Casos com CCS antes ≥ lim', i), c('Casos com CCS antes', i)), "0%", True)
prow("tsp", "Tratamentos sem produto informado", "trat.", lambda i: f"={c('Casos sem produto',i)}", "0", True)
ws.cell(row + 1, 1, "Estimativas: a CCS do tanque aqui é calculada pelas vacas do controle; não substitui a CCS "
        "real do laticínio (ainda não fornecida). Primeiro controle sem classificação (sem controle anterior).").font = font(
    italic=True, size=8, color="5E6B63")
widths(ws, {"A": 62, "B": 11, **{col: 9 for col in CL}})
ws.freeze_panes = "C5"
# destaque do controle selecionado
ws.conditional_formatting.add(f"C4:{CL[-1]}{row}", FormulaRule(
    formula=[f"COLUMN()-2=CISEL"], fill=PatternFill("solid", fgColor="DCE9F2")))

# ---------------------------------------------------------------- Lotes
ws = S["Lotes"]
title(ws, "Comparativo sanitário entre os lotes – controle selecionado",
      "O lote desviado (tratamento) entra na tabela mas fica fora da CCS estimada do tanque e do ranking.")
ws["A3"] = "=\"Controle: \"&Parâmetros!B%d" % (5 + len(params) + 1)
ws["A3"].font = font(bold=True)
lcols = ["Lote", "Vacas", "Prod. média", "CCS média", "CCS ponderada (estimativa)", "% ≥ limite",
         "Sadias", "Novas inf.", "Crônicas", "Recuperadas", "Pós-parto baixa", "Pós-parto elevada",
         "Sem histórico", "Taxa NI", "Recorrentes", "Mastite clínica no período", "% do tanque (sem desv.)"]
hdr(ws, 5, lcols)
LROW0 = 6
B_, D_, H_, I_, O_, T_, R_ = D("B"), D("D"), D("H"), D("I"), D("O"), D("T"), D("R")
for j, lt in enumerate(LOTES):
    r = LROW0 + j
    ws.cell(r, 1, lt)
    ws.cell(r, 2, f"=COUNTIFS({B_},CISEL,{D_},A{r})")
    ws.cell(r, 3, f'=IF(B{r}=0,"",SUMIFS({H_},{B_},CISEL,{D_},A{r})/B{r})')
    ws.cell(r, 4, f'=IF(B{r}=0,"",SUMIFS({I_},{B_},CISEL,{D_},A{r})/B{r})')
    ws.cell(r, 5, f'=IF(B{r}=0,"",SUMIFS({T_},{B_},CISEL,{D_},A{r})/SUMIFS({H_},{B_},CISEL,{D_},A{r}))')
    ws.cell(r, 6, f'=IF(B{r}=0,"",COUNTIFS({B_},CISEL,{D_},A{r},{I_},">="&LIM)/B{r})')
    for k, cl in enumerate(["S", "N", "C", "R", "PB", "PE", "H"]):
        ws.cell(r, 7 + k, f'=COUNTIFS({B_},CISEL,{D_},A{r},{O_},"{cl}")')
    ws.cell(r, 14, f'=IF(G{r}+H{r}=0,"",H{r}/(G{r}+H{r}))')
    ws.cell(r, 15, f"=COUNTIFS({B_},CISEL,{D_},A{r},{R_},1)")
    ws.cell(r, 16, f"=COUNTIFS({M('N')},CISEL,{M('M')},A{r})")
    ws.cell(r, 17, f'=IF(A{r}=DESVIO,"desviado",SUMIFS({T_},{B_},CISEL,{D_},A{r})/{csel("Σ prod×CCS sem desv.")})')
    for col, fm in ((3, "0.0"), (4, "0"), (5, "0"), (6, "0.0%"), (14, "0.0%"), (17, "0.0%")):
        ws.cell(r, col).number_format = fm
    for col in range(1, 18):
        ws.cell(r, col).border = BOX
LLAST = LROW0 + len(LOTES) - 1
r = LLAST + 1
ws.cell(r, 1, "Total").font = font(bold=True)
for col in [2] + list(range(7, 14)) + [15, 16]:
    ws.cell(r, col, f"=SUM({L(col)}{LROW0}:{L(col)}{LLAST})").font = font(bold=True)
ws.conditional_formatting.add(f"F{LROW0}:F{LLAST}", CellIsRule(operator="greaterThanOrEqual", formula=["0.3"],
                                                                 font=Font(name=F, color="C0392B", bold=True)))
# evolução por lote
r0 = LLAST + 4
ws.cell(r0 - 1, 1, "Evolução por lote: % de vacas com CCS ≥ limite (vazio = lote sem vacas no controle)").font = font(bold=True)
hdr(ws, r0, ["Lote"] + [f"=Controles!O{2+i}" for i in range(NC)])
for j, lt in enumerate(LOTES):
    r = r0 + 1 + j
    ws.cell(r, 1, lt)
    for i in range(NC):
        ws.cell(r, 2 + i, f'=IF(COUNTIFS({B_},{i+1},{D_},$A{r})=0,"",COUNTIFS({B_},{i+1},{D_},$A{r},{I_},">="&LIM)'
                          f'/COUNTIFS({B_},{i+1},{D_},$A{r}))').number_format = "0%"
r1 = r0 + len(LOTES) + 3
ws.cell(r1 - 1, 1, "Evolução por lote: novas infecções (nº)").font = font(bold=True)
hdr(ws, r1, ["Lote"] + [f"=Controles!O{2+i}" for i in range(NC)])
for j, lt in enumerate(LOTES):
    r = r1 + 1 + j
    ws.cell(r, 1, lt)
    for i in range(NC):
        ws.cell(r, 2 + i, f'=IF(COUNTIFS({B_},{i+1},{D_},$A{r})=0,"",COUNTIFS({B_},{i+1},{D_},$A{r},{O_},"N"))')
r2 = r1 + len(LOTES) + 3
ws.cell(r2 - 1, 1, "Evolução por lote: nº de vacas").font = font(bold=True)
hdr(ws, r2, ["Lote"] + [f"=Controles!O{2+i}" for i in range(NC)])
for j, lt in enumerate(LOTES):
    r = r2 + 1 + j
    ws.cell(r, 1, lt)
    for i in range(NC):
        ws.cell(r, 2 + i, f'=IF(COUNTIFS({B_},{i+1},{D_},$A{r})=0,"",COUNTIFS({B_},{i+1},{D_},$A{r}))')
widths(ws, {L(i): 10 for i in range(1, 18)})
ws.column_dimensions["A"].width = 12
ws.row_dimensions[5].height = 42

# ---------------------------------------------------------------- Histórico
ws = S["Histórico"]
title(ws, "Histórico por vaca – CCS em cada controle e sinalizações",
      "Sinalizações até o controle selecionado. Prioritária = elevada em todos (≥ 4 registros), ≥ 3 seguidos, recorrente no controle ou ≥ 3 tratamentos de mastite.")
hcols = (["Vaca"] + [f"=Controles!O{2+i}" for i in range(NC)] +
         ["Nº registros", "Nº com CCS ≥ limite", "Elevada em todos (≥ 4 reg.)", "Lote no controle",
          "Classe no controle", "Seguidos elevada", "Recorrências (total)", "Tratamentos de mastite",
          "Prioritária"])
hdr(ws, 4, hcols)
KEY = D("AD")
for j, v in enumerate(vacas):
    r = 5 + j
    ws.cell(r, 1, v)
    for i in range(NC):
        k = f"($A{r}*100+{i+1})"
        ws.cell(r, 2 + i, f'=IFERROR(IF(INDEX({KEY},MATCH({k},{KEY},1))={k},INDEX({D("I")},MATCH({k},{KEY},1)),""),"")')
    b = 2 + NC
    ws.cell(r, b, f"=COUNTIFS({D('A')},A{r},{D('B')},\"<=\"&CISEL)")
    ws.cell(r, b + 1, f"=COUNTIFS({D('A')},A{r},{D('B')},\"<=\"&CISEL,{D('J')},1)")
    ws.cell(r, b + 2, f'=IF(AND({L(b)}{r}>=4,{L(b+1)}{r}={L(b)}{r}),"sim","")')
    ksel = f"($A{r}*100+CISEL)"
    rowm = f"MATCH({ksel},{KEY},1)"
    ok = f"INDEX({KEY},{rowm})={ksel}"
    ws.cell(r, b + 3, f'=IFERROR(IF({ok},INDEX({D("D")},{rowm}),""),"")')
    ws.cell(r, b + 4, f'=IFERROR(IF({ok},INDEX({D("O")},{rowm}),""),"")')
    ws.cell(r, b + 5, f'=IFERROR(IF({ok},INDEX({D("Q")},{rowm}),""),"")')
    ws.cell(r, b + 6, f"=COUNTIFS({D('A')},A{r},{D('B')},\"<=\"&CISEL,{D('R')},1)")
    ws.cell(r, b + 7, f"=COUNTIFS({M('A')},A{r},{M('B')},\"<=\"&INDEX({CTR},CISEL))")
    cls_c, seq_c = L(b + 4), L(b + 5)
    ws.cell(r, b + 8, f'=IF({L(b+3)}{r}="","",IF(OR({L(b+2)}{r}="sim",N({seq_c}{r})>=3,'
                      f'AND({cls_c}{r}="N",IFERROR(INDEX({D("R")},{rowm}),0)=1),{L(b+7)}{r}>=3),"PRIORITÁRIA",""))')
widths(ws, {"A": 8, **{L(2 + i): 7 for i in range(NC)}, **{L(2 + NC + k): 11 for k in range(9)}})
ws.row_dimensions[4].height = 40
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A4:{L(1 + NC + 9)}{4 + NV}"
ws.conditional_formatting.add(f"B5:{L(1+NC)}{4+NV}", FormulaRule(
    formula=[f'AND(ISNUMBER(B5),B5>=LIM)'], font=Font(name=F, color="C0392B", bold=True)))
HFLAG = L(2 + NC + 8)

# ---------------------------------------------------------------- Ficha da vaca
ws = S["Ficha da vaca"]
title(ws, "Ficha da vaca", "Digite o número da vaca na célula amarela.")
ws["B3"] = "Vaca:"
ws["B3"].font = font(bold=True)
ws["C3"] = 2570
ws["C3"].fill = INFILL
ws["C3"].font = font(bold=True, size=12, color="0000FF")
ws["E3"] = f'=IF(COUNTIF({D("A")},C3)=0,"Vaca não encontrada nos controles",COUNTIF({D("A")},C3)&" registros de CCS · "&COUNTIF({M("A")},C3)&" tratamentos de mastite")'
hdr(ws, 5, ["Controle", "Data", "Lote", "Parto", "DEL", "Produção", "CCS", "Classe", "Seguidos elev.", "Recorrente", "Queda prod."])
for i in range(NC):
    r = 6 + i
    k = f"($C$3*100+{i+1})"
    rowm = f"MATCH({k},{KEY},1)"
    ok = f"INDEX({KEY},{rowm})={k}"
    ws.cell(r, 1, i + 1)
    ws.cell(r, 2, f"=INDEX({CTR},A{r})").number_format = "DD/MM/YYYY"
    for j, col in enumerate(["D", "E", "F", "H", "I", "O", "Q", "R", "S"]):
        ws.cell(r, 3 + j, f'=IFERROR(IF({ok},INDEX({D(col)},{rowm}),""),"")')
    ws.cell(r, 4).number_format = "DD/MM/YYYY"
    ws.cell(r, 7).number_format = "#,##0"
ws.conditional_formatting.add(f"G6:G{5+NC}", FormulaRule(formula=['AND(ISNUMBER(G6),G6>=LIM)'],
                                                        font=Font(name=F, color="C0392B", bold=True)))
t0 = 7 + NC
ws.cell(t0, 1, "Tratamentos de mastite clínica (CCS antes = último controle até o início; CCS depois = 1º controle após o retorno)").font = font(bold=True)
hdr(ws, t0 + 1, ["#", "Início", "Produto", "Quartos", "Data final", "Carência", "Retorno", "CCS antes", "CCS depois", "Retratamento", "Lote"])
for k in range(1, 31):
    r = t0 + 1 + k
    ws.cell(r, 1, k)
    rowm = f"MATCH($A{r},{M('R')},0)"
    for j, col in enumerate(["B", "C", "D", "E", "F", "G", "I", "K", "L", "M"]):
        ws.cell(r, 2 + j, f'=IFERROR(INDEX({M(col)},{rowm}),"")')
    for j in (2, 5, 7):
        ws.cell(r, j).number_format = "DD/MM/YYYY"
widths(ws, {"A": 9, "B": 12, "C": 14, "D": 12, "E": 10, "F": 10, "G": 12, "H": 10, "I": 10, "J": 11, "K": 10})

# ---------------------------------------------------------------- Ranking
ws = S["Ranking"]
title(ws, "Ranking do controle selecionado",
      "Top 50 maiores CCS (todas as vacas) · Top 20 maiores contribuições para a CCS estimada do tanque sem o lote desviado · simulação de descarte do Top N")
ws["A3"] = "=\"Controle: \"&Parâmetros!B%d" % (5 + len(params) + 1)
ws["A3"].font = font(bold=True)
rcols = ["#", "Vaca", "Lote", "DEL", "Produção", "CCS", "CCS anterior", "Classe", "Seguidos elev.",
         "Controles elev.", "Trat. mastite"]
ws["A5"] = "50 maiores CCS"
ws["A5"].font = font(bold=True, color=AZUL)
hdr(ws, 6, rcols)
RK_CCS0 = 7


def rank_row(r, k, keycol):
    ws.cell(r, 1, k)
    ln = f"(MATCH(LARGE({D(keycol)},$A{r}),{D(keycol)},0)+1)"
    ws.cell(r, 13, f"=IFERROR({ln},\"\")")  # linha em Dados (auxiliar)
    lr = f"$M{r}"
    for j, col in enumerate(["A", "D", "F", "H", "I", "N", "O", "Q"]):
        ws.cell(r, 2 + j, f'=IF({lr}="","",INDEX(Dados!${col}:${col},{lr}))')
    ws.cell(r, 10, f'=IF({lr}="","",COUNTIFS({D("A")},B{r},{D("B")},"<="&CISEL,{D("J")},1)&" de "&COUNTIFS({D("A")},B{r},{D("B")},"<="&CISEL))')
    ws.cell(r, 11, f'=IF({lr}="","",COUNTIFS({M("A")},B{r},{M("B")},"<="&INDEX({CTR},CISEL)))')
    ws.cell(r, 5).number_format = "0.0"
    ws.cell(r, 6).number_format = "#,##0"
    ws.cell(r, 7).number_format = "#,##0"
    for col in range(1, 12):
        ws.cell(r, col).border = BOX


for k in range(1, 51):
    rank_row(RK_CCS0 + k - 1, k, "X")
r0 = RK_CCS0 + 50 + 2
ws.cell(r0, 1, "20 maiores contribuições para a CCS estimada do tanque (sem o lote desviado)").font = font(bold=True, color=AZUL)
hdr(ws, r0 + 1, rcols + ["", "Prod × CCS", "% do tanque", "% acumulado"])
ws.cell(r0 + 1, 12).fill = PatternFill()
RK_CT0 = r0 + 2
TOT = csel("Σ prod×CCS sem desv.")
for k in range(1, 21):
    r = RK_CT0 + k - 1
    rank_row(r, k, "Y")
    ws.cell(r, 14, f'=IF($M{r}="","",INDEX(Dados!$T:$T,$M{r}))').number_format = "#,##0"
    ws.cell(r, 15, f'=IF($M{r}="","",N{r}/{TOT})').number_format = "0.0%"
    ws.cell(r, 16, f'=IF($M{r}="","",SUM($N${RK_CT0}:N{r})/{TOT})').number_format = "0.0%"
s0 = RK_CT0 + 21
ws.cell(s0, 1, "Simulação de descarte (ou desvio do leite) do Top N em contribuição").font = font(bold=True, color=AZUL)
sim = [
    ("Top N (Parâmetros)", "=TOPN", "0"),
    ("CCS estimada do tanque sem o lote desviado – atual (estimativa)",
     f"={csel('Σ prod×CCS sem desv.')}/{csel('Σ prod sem desv.')}", "0"),
    ("Produção retirada (kg)", f"=SUMIF(Calc!$A${RK1}:$A${RK2},\"<=\"&TOPN,Calc!$D${RK1}:$D${RK2})", "#,##0.0"),
    ("% da produção sem o lote desviado", f"=B{s0+3}/{csel('Σ prod sem desv.')}", "0.0%"),
    ("CCS estimada sem o Top N (estimativa)",
     f"=({csel('Σ prod×CCS sem desv.')}-SUMIF(Calc!$A${RK1}:$A${RK2},\"<=\"&TOPN,Calc!$E${RK1}:$E${RK2}))"
     f"/({csel('Σ prod sem desv.')}-B{s0+3})", "0"),
    ("Redução estimada (mil/mL)", f"=B{s0+2}-B{s0+5}", "0"),
]
for j, (t, f, fm) in enumerate(sim):
    ws.cell(s0 + 1 + j, 1, t)
    c_ = ws.cell(s0 + 1 + j, 2, f)
    c_.number_format = fm
    c_.font = font(bold=True)
ws.cell(s0 + 8, 1, "Simulação aritmética: não considera reposição de vacas nem mudança de produção. "
        "Decisões de tratamento, secagem ou descarte são sempre tomadas com o veterinário responsável.").font = font(italic=True, size=8)
widths(ws, {"A": 5, "B": 8, "C": 6, "D": 6, "E": 9, "F": 8, "G": 9, "H": 8, "I": 9, "J": 11, "K": 9, "L": 2, "M": 8, "N": 11, "O": 10, "P": 11})
ws.column_dimensions["M"].hidden = True
# a coluna A da simulação é longa: texto sobrepõe; ok
SIMROW = s0 + 1

# ---------------------------------------------------------------- Dados gráficos (oculta)
ws = S["Dados gráficos"]
ws["A1"] = "Séries dos gráficos (só os controles carregados)"
gcols = ["Controle", "CCS est. tanque sem desv.", "Média aritmética sem desv.", "Média geométrica sem desv.",
         "Sadia", "Recuperada", "Nova infecção", "Crônica", "Rótulo com n", "< 200", "200–499", "500–999", "≥ 1.000",
         "Novas infecções", "Recuperadas", "Taxa NI", "Tratamentos mastite", "Trat./100 vacas/30 d",
         "% ≥ limite", "Produção média", "Crônicas", "Recorrentes"]
hdr(ws, 2, gcols)
for i in range(NC):
    r = 3 + i
    ws.cell(r, 1, f"=Controles!O{2+i}")
    for j, k in enumerate(["i1", "arit", "geo"]):
        ws.cell(r, 2 + j, f"=ROUND(Painel!{CL[i]}{PROW[k]},0)")
    for j, k in enumerate(["S", "R", "N", "C"]):
        ws.cell(r, 5 + j, f"=IF(Painel!{CL[i]}{PROW['base']}=\"\",\"\",ROUND({c(k,i)}/Painel!{CL[i]}{PROW['base']},2))"
                if i else '=""')
    ws.cell(r, 9, f'=Controles!O{2+i}&" n="&Painel!{CL[i]}{PROW["base"]}')
    for j, k in enumerate(["f1", "f2", "f3", "f4"]):
        ws.cell(r, 10 + j, f"=Painel!{CL[i]}{PROW[k]}")
    ws.cell(r, 14, f"=Painel!{CL[i]}{PROW['nN']}")
    ws.cell(r, 15, f"=Painel!{CL[i]}{PROW['nR']}")
    ws.cell(r, 16, f"=Painel!{CL[i]}{PROW['tni']}")
    ws.cell(r, 17, f"=Painel!{CL[i]}{PROW['trat']}")
    ws.cell(r, 18, f"=IF(Painel!{CL[i]}{PROW['t100']}=\"\",\"\",ROUND(Painel!{CL[i]}{PROW['t100']},1))")
    ws.cell(r, 19, f"=ROUND(Painel!{CL[i]}{PROW['i6']},3)")
    ws.cell(r, 20, f"=ROUND(Painel!{CL[i]}{PROW['i9']},1)")
    ws.cell(r, 21, f"=Painel!{CL[i]}{PROW['nC']}")
    ws.cell(r, 22, f"=Painel!{CL[i]}{PROW['rec']}")
    for j in (5, 6, 7, 8, 16, 19):
        ws.cell(r, j).number_format = "0%"
GL = 2 + NC  # última linha de dados dos gráficos
ws.sheet_state = "hidden"
GD = ws


def labels(ch, pos=None, pct=False, fmt=None):
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    for a in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
        setattr(ch.dataLabels, a, False)
    if pos:
        ch.dataLabels.position = pos
    if fmt:
        ch.dataLabels.numFmt = fmt


def size(ch):
    ch.width, ch.height = 18, 11


def mk_line(ttl, cols, colors, first=3, ytitle="mil céls/mL"):
    ch = LineChart()
    ch.title = ttl
    ch.y_axis.title = ytitle
    for col, colr in zip(cols, colors):
        ch.add_data(Reference(GD, min_col=col, min_row=2, max_row=GL), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties.line.solidFill = colr
        s.graphicalProperties.line.width = 28000
        s.marker.symbol = "circle"
        s.marker.size = 6
        s.marker.graphicalProperties = GraphicalProperties(solidFill=colr)
        s.marker.graphicalProperties.line.solidFill = colr
        s.smooth = False
    ch.set_categories(Reference(GD, min_col=1, min_row=first, max_row=GL))
    labels(ch, "t")
    ch.legend.position = "b"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.y_axis.majorGridlines = None
    size(ch)
    return ch


def mk_bar(ttl, cols, colors, grouping="clustered", first=3, cats_col=1, pct=False, ytitle="vacas"):
    ch = BarChart()
    ch.type = "col"
    ch.grouping = grouping
    if grouping != "clustered":
        ch.overlap = 100
    ch.title = ttl
    ch.y_axis.title = ytitle
    for col, colr in zip(cols, colors):
        ch.add_data(Reference(GD, min_col=col, min_row=2 if first == 3 else first - 1, max_row=GL), titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties = GraphicalProperties(solidFill=colr)
        s.graphicalProperties.line.solidFill = "FFFFFF" if colr != "FFFFFF" else "9AA39D"
    ch.set_categories(Reference(GD, min_col=cats_col, min_row=first, max_row=GL))
    labels(ch, "ctr" if grouping != "clustered" else "outEnd", fmt="0%" if pct else None)
    ch.gapWidth = 60
    ch.legend.position = "b"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.y_axis.majorGridlines = None
    if pct:
        ch.y_axis.number_format = "0%"
        ch.y_axis.scaling.max = 1
    size(ch)
    return ch


def charts():
    """Recria os gráficos (cada aba precisa de objetos próprios)."""
    c1 = mk_line("Evolução da CCS do rebanho sem o lote 10 (estimativa)", [2, 3, 4], [AZUL, "B7791F", "2E8B57"])
    # classificação: só a partir do 2º controle (o 1º não tem anterior)
    c2 = BarChart()
    c2.type, c2.grouping, c2.overlap = "col", "percentStacked", 100
    c2.title = "Classificação sanitária (% das vacas na base)"
    for col, k in zip([5, 6, 7, 8], ["S", "R", "N", "C"]):
        ref = Reference(GD, min_col=col, min_row=4, max_row=GL)
        from openpyxl.chart import Series
        s = Series(ref, title=NOME[k])
        s.graphicalProperties = GraphicalProperties(solidFill=COR[k])
        s.graphicalProperties.line.solidFill = "FFFFFF"
        c2.series.append(s)
    c2.set_categories(Reference(GD, min_col=9, min_row=4, max_row=GL))
    labels(c2, "ctr", fmt="0%")
    c2.gapWidth = 50
    c2.legend.position = "b"
    c2.y_axis.number_format = "0%"
    c2.y_axis.delete = False
    c2.x_axis.delete = False
    c2.y_axis.majorGridlines = None
    size(c2)
    c3 = mk_bar("Vacas por faixa de CCS", [10, 11, 12, 13], ["2E8B57", "E8C547", "E8833A", "C0392B"], "stacked")
    c4 = mk_bar("Novas infecções, crônicas e recuperadas (nº de vacas)", [14, 21, 15], ["1B3A6B", "C0392B", "9AA39D"],
                first=4)
    c5 = mk_bar("Mastite clínica: tratamentos por 100 vacas a cada 30 dias", [18], ["D9713C"], first=4, ytitle="trat.")
    c6 = mk_line("% de vacas com CCS ≥ 200", [19], ["C0392B"], ytitle="%")
    c6.y_axis.number_format = "0%"
    c6.dataLabels.numFmt = "0.0%"
    return [c1, c2, c3, c4, c5, c6]


# ---------------------------------------------------------------- Gráficos
ws = S["Gráficos"]
title(ws, "Gráficos – controles carregados", "Dois gráficos por página A4, 18 × 11 cm. Sem o lote 10 na evolução da CCS.")
chs = charts()
for i, ch in enumerate(chs):
    ws.add_chart(ch, f"A{4 + i * 23 + (i // 2) * 2}")
for i in (1, 2):
    ws.row_breaks.append(Break(id=3 + i * 48 - 2))
ws.print_area = f"A1:K{4 + 6 * 23 + 8}"

# ---------------------------------------------------------------- Impressão
ws = S["Impressão"]
for i in range(1, 13):
    ws.column_dimensions[L(i)].width = 7.4
ws.column_dimensions["A"].width = 5
selrow = 5 + len(params) + 1
ws["A1"] = "Qualidade do Leite · ARKAFLA"
ws["A1"].font = font(bold=True, size=14, color=AZUL)
ws["A2"] = (f'="Controle de "&Parâmetros!B{selrow}&" · REB "&INDEX(Controles!$C$2:$C${NC+1},CISEL)&" · "&'
            f'TEXT({csel("Vacas")},"#.##0")&" vacas · "&IF(CISEL>1,INDEX(Controles!$H$2:$H${NC+1},CISEL)&" dias desde o controle anterior","primeiro controle")')
ws["A2"].font = font(size=9, color="5E6B63")
ws["A4"] = "1. Indicadores do controle"
ws["A4"].font = font(bold=True, color=AZUL)
hdr(ws, 5, ["Indicador", "", "", "", "", "", "Atual", "", "Anterior", "", "Variação", "Nº vacas"])
for a, b in (("A5", "F5"), ("G5", "H5"), ("I5", "J5")):
    ws.merge_cells(f"{a}:{b}")
ind = [("i1", "CCS est. tanque sem lote 10 (estimativa)", "0", "abs", None),
       ("apc", '"Média CCS tanque" APCBRH (estimativa)', "0", "abs", None),
       ("i9", "Produção média (kg/vaca/dia)", "0.0", "abs", "n"),
       ("i6", "% vacas CCS ≥ 200", "0.0%", "pp", "nel"),
       ("i7", "% vacas CCS ≥ 500", "0.0%", "pp", None),
       ("i8", "% vacas CCS ≥ 1.000", "0.0%", "pp", "f4"),
       ("i2", "% sadias", "0.0%", "pp", "nS"),
       ("i3", "% novas infecções", "0.0%", "pp", "nN"),
       ("i4", "% crônicas", "0.0%", "pp", "nC"),
       ("i5", "% recuperadas", "0.0%", "pp", "nR"),
       ("tni", "Taxa de novas infecções", "0.0%", "pp", "nN"),
       ("tni30", "Taxa NI por 30 dias (estimativa)", "0.0%", "pp", None),
       ("perm", "Permanência", "0.0%", "pp", "nC"),
       ("recp", "Recorrência", "0.0%", "pp", "rec"),
       ("i10", "Elevada ≥ 2 controles seguidos", "0", "abs", "i10"),
       ("seq3", "Elevada ≥ 3 controles seguidos", "0", "abs", "seq3"),
       ("queda", "Elevada + queda de produção", "0", "abs", "queda"),
       ("cura", "Cura no período seco", "0.0%", "pp", "curan"),
       ("nis", "Nova infecção no período seco", "0.0%", "pp", "nisn"),
       ("PE", "Pós-parto elevada (fora da base)", "0", "abs", "PE"),
       ("H", "Sem histórico (fora da base)", "0", "abs", "H"),
       ("base", "Vacas na base da classificação", "#,##0", "abs", "base"),
       ("trat", "Tratamentos de mastite no período", "0", "abs", "trat"),
       ("t100", "Trat. por 100 vacas a cada 30 dias", "0.0", "abs", None)]
PR = f"Painel!$C${{r}}:${CL[-1]}${{r}}"
for j, (k, t, fm, var, nk) in enumerate(ind):
    r = 6 + j
    pr = PROW[k]
    rng = f"Painel!$C${pr}:${CL[-1]}${pr}"
    ws.merge_cells(f"A{r}:F{r}")
    ws.merge_cells(f"G{r}:H{r}")
    ws.merge_cells(f"I{r}:J{r}")
    ws.cell(r, 1, t)
    ws.cell(r, 7, f"=INDEX({rng},CISEL)").number_format = fm
    ws.cell(r, 9, f'=IF(CISEL>1,INDEX({rng},CISEL-1),"")').number_format = fm
    if var == "pp":
        ws.cell(r, 11, f'=IFERROR(IF(OR(G{r}="",I{r}=""),"",(G{r}-I{r})*100),"")').number_format = '+0.0" p.p.";-0.0" p.p.";0.0" p.p."'
    else:
        ws.cell(r, 11, f'=IFERROR(IF(OR(G{r}="",I{r}=""),"",G{r}-I{r}),"")').number_format = "+" + fm.replace("#,##", "") + ";-" + fm.replace("#,##", "")
    if nk:
        ws.cell(r, 12, f"=INDEX(Painel!$C${PROW[nk]}:${CL[-1]}${PROW[nk]},CISEL)").number_format = "#,##0"
    for col in range(1, 13):
        ws.cell(r, col).border = BOX
        ws.cell(r, col).font = font(size=9)
r = 6 + len(ind) + 1
ws.cell(r, 1, "2. Lotes no controle").font = font(bold=True, color=AZUL)
hdr(ws, r + 1, ["Lote", "Vacas", "Prod.", "CCS méd.", "% ≥ 200", "Sadias", "Novas inf.", "Crôn.", "Recup.",
                "Pós-p. elev.", "Mastite", "% tanque"])
lmap = [1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 16, 17]
for j in range(len(LOTES) + 1):
    rr = r + 2 + j
    for col, src in enumerate(lmap, 1):
        cc = ws.cell(rr, col, f'=IF(Lotes!$B{LROW0+j}=0,"",Lotes!{L(src)}{LROW0+j})' if j < len(LOTES)
                     else (f"=Lotes!{L(src)}{LROW0+j}" if src in (2, 7, 8, 9, 10, 12, 16) else ("Total" if col == 1 else None)))
        cc.font = font(size=9)
        cc.border = BOX
        cc.number_format = {3: "0.0", 4: "0", 5: "0%", 12: "0%"}.get(col, "0")
P1END = r + 2 + len(LOTES)
ws.cell(P1END + 1, 1, "Estimativas calculadas pelas vacas do controle; a CCS real do tanque (laticínio) ainda não foi fornecida.").font = font(italic=True, size=8)
ws.row_breaks.append(Break(id=P1END + 2))
# página 2: gráficos
p2 = P1END + 3
ws.cell(p2, 1, "3. Evolução da CCS sem o lote 10 e classificação sanitária").font = font(bold=True, color=AZUL)
ch = charts()
ch[0].width, ch[0].height = 18, 11
ch[1].width, ch[1].height = 18, 11
ws.add_chart(ch[0], f"A{p2 + 1}")
ws.add_chart(ch[1], f"A{p2 + 24}")
ws.cell(p2 + 46, 1, "Classificação: base = vacas com CCS neste controle e no anterior; recém-paridas e vacas sem registro anterior ficam fora da base.").font = font(italic=True, size=8)
ws.row_breaks.append(Break(id=p2 + 47))
# página 3: top 50 CCS
p3 = p2 + 48
ws.cell(p3, 1, "4. Ranking: 50 maiores CCS no controle").font = font(bold=True, color=AZUL)
hdr(ws, p3 + 1, ["#", "Vaca", "Lote", "DEL", "Prod.", "CCS", "CCS ant.", "Classe", "", "Seguidos", "Contr. elev.", "Trat."])
ws.merge_cells(f"H{p3+1}:I{p3+1}")
for k in range(50):
    rr = p3 + 2 + k
    src = RK_CCS0 + k
    for col, sc in zip([1, 2, 3, 4, 5, 6, 7, 8, 10, 11, 12], ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]):
        cc = ws.cell(rr, col, f"=Ranking!{sc}{src}" if sc != "H" else
                     f'=IFERROR(CHOOSE(MATCH(Ranking!H{src},{{"S","N","C","R","PB","PE","H"}},0),"Sadia","Nova infecção","Crônica","Recuperada","Pós-parto baixa","Pós-parto elevada","Sem histórico"),"")')
        cc.font = font(size=8.5)
        cc.border = BOX
        cc.number_format = {5: "0.0", 6: "#,##0", 7: "#,##0"}.get(col, "General")
    ws.merge_cells(f"H{rr}:I{rr}")
ws.row_breaks.append(Break(id=p3 + 52))
# página 4: top 20 contribuições + simulação
p4 = p3 + 53
ws.cell(p4, 1, "5. Ranking: 20 maiores contribuições para a CCS estimada do tanque (sem o lote 10)").font = font(bold=True, color=AZUL)
hdr(ws, p4 + 1, ["#", "Vaca", "Lote", "DEL", "Prod.", "CCS", "Classe", "", "Contr. elev.", "Trat.", "% tanque", "% acum."])
ws.merge_cells(f"G{p4+1}:H{p4+1}")
for k in range(20):
    rr = p4 + 2 + k
    src = RK_CT0 + k
    for col, sc in zip([1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12], ["A", "B", "C", "D", "E", "F", "H", "J", "K", "O", "P"]):
        cc = ws.cell(rr, col, f"=Ranking!{sc}{src}" if sc != "H" else
                     f'=IFERROR(CHOOSE(MATCH(Ranking!H{src},{{"S","N","C","R","PB","PE","H"}},0),"Sadia","Nova infecção","Crônica","Recuperada","Pós-parto baixa","Pós-parto elevada","Sem histórico"),"")')
        cc.font = font(size=8.5)
        cc.border = BOX
        cc.number_format = {5: "0.0", 6: "#,##0", 11: "0.0%", 12: "0.0%"}.get(col, "General")
    ws.merge_cells(f"G{rr}:H{rr}")
rr = p4 + 23
ws.cell(rr, 1, "Simulação de descarte do Top N").font = font(bold=True, color=AZUL)
for j, (t, _, fm) in enumerate(sim):
    ws.merge_cells(f"A{rr+1+j}:H{rr+1+j}")
    ws.cell(rr + 1 + j, 1, t).font = font(size=9)
    cc = ws.cell(rr + 1 + j, 9, f"=Ranking!B{SIMROW + j}")
    cc.number_format = fm
    cc.font = font(size=9, bold=True)
ws.cell(rr + 8, 1, "Decisões de tratamento, secagem ou descarte são sempre tomadas com o veterinário responsável.").font = font(italic=True, size=8)
ws.print_area = f"A1:L{rr + 8}"
for w in (ws, S["Gráficos"]):
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.orientation = "portrait"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_margins.left = w.page_margins.right = 0.47
    w.page_margins.top = w.page_margins.bottom = 0.47
    w.print_options.horizontalCentered = True

# ---------------------------------------------------------------- Cultura
ws = S["Cultura"]
title(ws, "Cultura microbiológica", "Ainda não fornecida. Preencher uma linha por amostra quando os resultados chegarem.")
hdr(ws, 4, ["Vaca", "Data da coleta", "Amostra (quarto ou composta)", "Quarto", "Agente", "Antibiograma – sensível",
            "Antibiograma – resistente", "Laboratório", "CCS no controle anterior", "Observação"])
for r in range(5, 105):
    ws.cell(r, 9, f'=IF(A{r}="","",IFERROR(INDEX({D("I")},MATCH(A{r}*100+_xlfn.MAXIFS({D("B")},{D("A")},A{r},{D("C")},"<="&B{r}),{KEY},0)),""))')
    ws.cell(r, 2).number_format = "DD/MM/YYYY"
widths(ws, {L(i): 16 for i in range(1, 11)})

# ---------------------------------------------------------------- Critérios
ws = S["Critérios"]
title(ws, "Critérios de classificação e fórmulas")
hdr(ws, 3, ["Situação", "Controle anterior", "Controle atual", "Cor", "Entra na base dos %?"])
crit = [("S", "< 200 mil", "< 200 mil", "sim"), ("N", "< 200 mil", "≥ 200 mil", "sim"),
        ("C", "≥ 200 mil", "≥ 200 mil", "sim"), ("R", "≥ 200 mil", "< 200 mil", "sim"),
        ("PB", "último da lactação anterior", "1º da nova lactação < 200", "não (avalia período seco)"),
        ("PE", "último da lactação anterior", "1º da nova lactação ≥ 200", "não (avalia período seco)"),
        ("H", "sem registro ou faltou ao controle anterior", "qualquer", "não; nunca conta como nova infecção")]
for j, (k, a, b, base) in enumerate(crit):
    r = 4 + j
    ws.cell(r, 1, NOME[k])
    ws.cell(r, 2, a)
    ws.cell(r, 3, b)
    ws.cell(r, 4).fill = PatternFill("solid", fgColor=COR[k])
    ws.cell(r, 4).border = Border(*(Side(style="thin", color="9AA39D"),) * 4)
    ws.cell(r, 5, base)
txt = [
    "Limite individual: 200 mil céls/mL (Parâmetros). Faixas: < 200, 200–499, 500–999, ≥ 1.000.",
    "Comparação sempre com o registro anterior da mesma vaca, na mesma lactação. Nova lactação = parto posterior à data do controle anterior.",
    "Base dos % de Sadia, Nova infecção, Crônica e Recuperada: vacas com CCS no controle atual e no anterior (consecutivos), sem recém-paridas e sem 'sem histórico'.",
    "Taxa de novas infecções = NI ÷ (Sadias + NI). Taxa por 30 dias = taxa × 30 ÷ dias do intervalo (estimativa; só para intervalos ≥ 28 dias).",
    "Permanência = Crônicas ÷ (Crônicas + Recuperadas). Recorrente = NI em vaca que já teve CCS ≥ 200 na mesma lactação. Recorrência = recorrentes ÷ NI.",
    "Controles seguidos com CCS elevada: soma 1 a cada controle consecutivo na mesma lactação; uma falta no controle zera a contagem.",
    "Cura no período seco = pós-parto baixa ÷ vacas que secaram elevadas; nova infecção no período seco = pós-parto elevada ÷ vacas que secaram baixas.",
    "Queda de produção: CCS elevada, produção ≤ 70% do registro anterior da mesma lactação e DEL ≥ 60.",
    "CCS estimada do tanque = Σ(produção × CCS) ÷ Σ produção. 'Sem o lote 10' exclui o lote de tratamento (leite desviado). É estimativa, não a CCS real do laticínio.",
    "Mastite clínica: retratamento = mesma vaca tratada de novo em até 21 dias do início anterior. CCS antes = último controle até o início; CCS depois = 1º controle após o retorno do leite. Lote do caso = lote no controle anterior ao caso.",
    "Tratamentos por 100 vacas a cada 30 dias = tratamentos iniciados no intervalo ÷ vacas avaliadas × 100 × 30 ÷ dias do intervalo.",
    "A escolha do produto não foi aleatória: não comparar eficácia de produtos. Redução da CCS não comprova cura microbiológica.",
]
for j, t in enumerate(txt):
    ws.cell(13 + j, 1, t)
widths(ws, {"A": 20, "B": 34, "C": 28, "D": 8, "E": 34})

# ---------------------------------------------------------------- Qualidade dos dados
ws = S["Qualidade dos dados"]
title(ws, "Qualidade dos dados e inconsistências")
hdr(ws, 3, ["Controle", "REB", "Vacas no relatório", "Linhas lidas", "Confere média APCBRH?", "Partos em branco / DEL 0",
            "CCS no teto (9.998)", "Partos alterados", "Entraram", "Saíram", "Lotes com < 10 vacas"])
for i in range(NC):
    r = 4 + i
    ws.cell(r, 1, f"=Controles!O{2+i}")
    ws.cell(r, 2, f"=Controles!C{2+i}")
    ws.cell(r, 3, f"=Controles!D{2+i}")
    ws.cell(r, 4, f"=Controles!I{2+i}")
    ws.cell(r, 5, f"=Controles!N{2+i}")
    ws.cell(r, 6, f"={c('Partos em branco',i)}")
    ws.cell(r, 7, f"={c('CCS no teto',i)}")
    ws.cell(r, 8, f"={c('Partos alterados',i)}")
    ws.cell(r, 9, f"=Painel!{CL[i]}{PROW['ent']}")
    ws.cell(r, 10, f"=Painel!{CL[i]}{PROW['sai']}")
    few = [lt for lt in LOTES]
    ws.cell(r, 11, "=" + "&".join(
        f'IF(AND(COUNTIFS({B_},{i+1},{D_},{lt})>0,COUNTIFS({B_},{i+1},{D_},{lt})<10),"{lt} ("&COUNTIFS({B_},{i+1},{D_},{lt})&") ","")'
        for lt in few))
r = 4 + NC
ws.cell(r, 1, "Total")
for col in (6, 7, 8):
    ws.cell(r, col, f"=SUM({L(col)}4:{L(col)}{r-1})")
alt = sorted({x["vaca"] for x in recs if x["parto_alt"]}, key=int)
r += 2
ws.cell(r, 1, "Controle leiteiro").font = font(bold=True, color=AZUL)
notes = [
    f"Partos alterados entre controles (mesma lactação, data de parto diferente): vacas {', '.join(alt)}.",
    "CCS 9.998 é provavelmente o teto de leitura do equipamento; mantida como informada.",
    "Lotes reorganizados em junho. Lotes 7, 13, 14 e 20 com poucas vacas.",
    "REB 47238 até 18/06/2026 e REB 70105 a partir de 17/07/2026: mesmo rebanho: mudança no código do produtor confirmada pela APCBRH (informado em 01/10/2026).",
    "O arquivo 'r22_janeiro.xlsx' corresponde ao controle de 02/02/2026 (referente a janeiro). Não há controle com data em janeiro.",
    "Em todos os controles a média ponderada recalculada bate com a 'MÉDIA CCS TANQUE' da APCBRH (coluna 'Confere?').",
]
for t in notes:
    r += 1
    ws.cell(r, 1, t)
r += 2
ws.cell(r, 1, "Registro de mastite clínica").font = font(bold=True, color=AZUL)
mnotes = [
    (f'="Tratamentos sem produto informado (\'Mastite clinica\'): "&SUM({M("O")})', None),
    (f'="Tratamentos sem quarto informado: "&SUM({M("P")})', None),
    (f'="Tratamentos de vacas que não aparecem nos controles: "&COUNTIF({M("S")},"NÃO")', None),
    ("11 tratamentos com datas inconsistentes no registro original.", None),
    ("1 linha duplicada (vaca 4687, 20/04/2026): mantida como no registro; conta como retratamento.", None),
    ("1 retorno digitado '297/26', recalculado como 29/07/2026.", None),
    ("Anotação 'MORTA' no quarto PE da vaca 6000.", None),
] + [(t, None) for t in mast.get("notas", [])]
for t, _ in mnotes:
    r += 1
    ws.cell(r, 1, t)
widths(ws, {L(i): 12 for i in range(1, 12)})
ws.column_dimensions["K"].width = 24

# ---------------------------------------------------------------- Leia-me
ws = S["Leia-me"]
title(ws, "ARKAFLA – Qualidade do Leite 2026", "Fazenda ARKAFLA (Castro/PR) · controle leiteiro APCBRH e registro de mastite clínica")
lm = [
    ("Conteúdo", ""),
    ("Painel", "10 indicadores do projeto e seções A (rebanho), B (glândula mamária) e C (mastite clínica), uma coluna por controle."),
    ("Impressão", "Folha A4 retrato em 4 páginas do controle selecionado (Parâmetros → controle analisado)."),
    ("Gráficos", "Seis gráficos, dois por página A4 (18 × 11 cm)."),
    ("Lotes", "Comparativo sanitário por lote no controle selecionado e evolução por lote."),
    ("Ranking", "50 maiores CCS, 20 maiores contribuições sem o lote 10 e simulação de descarte do Top N."),
    ("Histórico", "Uma linha por vaca, CCS por controle e sinalizações (filtro na coluna 'Prioritária')."),
    ("Ficha da vaca", "Digite a vaca: CCS por controle e tratamentos de mastite."),
    ("Mastite clínica", "Um tratamento por linha, com CCS antes e depois, retratamento e lote do caso."),
    ("Cultura", "Modelo para os resultados de cultura (ainda não fornecidos)."),
    ("Critérios / Qualidade dos dados / Parâmetros / Controles", "Regras, inconsistências, parâmetros e metadados de cada relatório."),
    ("Dados", "Todos os controles empilhados, ordenados por vaca e nº do controle; classificação por fórmulas."),
    ("Calc / Dados gráficos", "Cálculos auxiliares (Dados gráficos fica oculta)."),
    ("", ""),
    ("Controles carregados", "=TEXT(MIN(Controles!B2:B%d),\"DD/MM/YYYY\")&\" a \"&TEXT(MAX(Controles!B2:B%d),\"DD/MM/YYYY\")&\" (\"&COUNT(Controles!A2:A%d)&\" controles)\"" % (NC + 1, NC + 1, NC + 1)),
    ("Como atualizar", "Anexar o novo Relatório 2.2 e o registro de mastite; o script scripts/gera_planilha.py refaz esta planilha."),
    ("Estimativas", "A CCS do tanque aqui é sempre estimativa pelas vacas do controle. A CCS real do laticínio ainda não foi fornecida."),
    ("Pendências", "Tabela de faixas do prêmio de CCS do laticínio; tratamentos sistêmicos de mastite (aba DOENÇAS GERAL) ainda fora dos indicadores."),
    ("Decisões", "Tratamento, secagem ou descarte: sempre com o veterinário responsável."),
]
for j, (a, b) in enumerate(lm):
    ws.cell(4 + j, 1, a).font = font(bold=True)
    ws.cell(4 + j, 2, b)
widths(ws, {"A": 34, "B": 120})

# fonte Arial em tudo o que ficou sem fonte explícita
for w in wb.worksheets:
    for row_ in w.iter_rows():
        for cell in row_:
            if cell.font is None or cell.font.name != F:
                f0 = cell.font
                cell.font = Font(name=F, size=f0.size or 10, bold=f0.bold, italic=f0.italic, color=f0.color)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("ok", OUT, "linhas Dados", NR, "tratamentos", NM, "vacas", NV)
