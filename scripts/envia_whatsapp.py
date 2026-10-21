"""Envia o PDF do monitoramento ARCAFLA e um resumo em texto pelo WhatsApp do usuário, via Z-API.

O WhatsApp do usuário fica conectado ao Z-API (QR Code, como no WhatsApp Web). Configuração por
variáveis de ambiente (nunca gravadas em arquivo):
    ZAPI_INSTANCE      ID da instância no painel do Z-API
    ZAPI_TOKEN         token da instância
    ZAPI_CLIENT_TOKEN  token de segurança da conta (opcional; obrigatório se ativado no painel)
    WPP_DESTINO        número(s) de destino com DDI e DDD, só dígitos, separados por vírgula
                       (ex.: 5543999998888)
O domínio api.z-api.io precisa estar liberado na rede do ambiente.

Uso:
    python3 scripts/envia_whatsapp.py [pdf] [planilha.xlsx]
Sem argumentos usa o PDF mais recente em monitoramento/ e a planilha Monitoramento_Diario_ARCAFLA.xlsx.
"""
import base64
import glob
import json
import os
import re
import sys
import urllib.request
from datetime import datetime

from openpyxl import load_workbook

PDF = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob("monitoramento/Monitoramento_Diario_ARCAFLA_*.pdf"))[-1]
XLSX = sys.argv[2] if len(sys.argv) > 2 else "monitoramento/Monitoramento_Diario_ARCAFLA.xlsx"
BASE = "https://api.z-api.io/instances/{inst}/token/{tok}"


def br(v, casas=0):
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def resumo():
    """Resumo curto do último dia, lido da planilha recalculada."""
    ws = load_workbook(XLSX, data_only=True)["Monitoramento"]
    cab = [ws.cell(4, j).value for j in range(1, ws.max_column + 1)]
    dcols = [j + 1 for j, v in enumerate(cab) if isinstance(v, datetime)]
    ult, ant = dcols[-1], (dcols[-2] if len(dcols) > 1 else None)
    lin = {str(ws.cell(r, 1).value).strip(): r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}

    def v(rot, col=ult):
        x = ws.cell(lin[rot], col).value if rot in lin else None
        if isinstance(x, str):
            m = re.match(r"\s*(-?[\d.]+)", x)
            return float(m[1]) if m else None
        return x

    def d(rot, casas=0):
        a, b = v(rot), v(rot, ant) if ant else None
        if not isinstance(a, (int, float)) or not isinstance(b, (int, float)):
            return ""
        x = a - b
        return f" ({'+' if x > 0 else ''}{br(x, casas)})" if abs(x) > 1e-9 else " (=)"

    dia = cab[ult - 1]
    t = [f"*ARCAFLA – Monitoramento {dia:%d/%m/%Y}*"]
    t.append(f"Vacas em lactação: {br(v('Vacas em lactação (lotes 1 a 6 e 10)'))}"
             f"{d('Vacas em lactação (lotes 1 a 6 e 10)')}")
    p, e = v("Produção média nas últimas 24 h (DP)"), v("Média de leite ECM (gordura e proteína do tanque)")
    if p:
        t.append(f"Produção média: {br(p, 1)} kg/vaca{d('Produção média nas últimas 24 h (DP)', 1)}"
                 + (f" · ECM {br(e, 1)}" if e else ""))
    t.append(f"Lote 10 (tratamento): {br(v('Animais em tratamento (lote 10)'))}"
             f"{d('Animais em tratamento (lote 10)')}")
    pt = v("Vacas que começaram a lactação (partos)")
    if isinstance(pt, (int, float)):
        t.append(f"Partos: {br(pt)}")
    g = v("% gestantes ÷ (adultas − DNB)")
    if isinstance(g, (int, float)):
        t.append(f"Gestantes: {br(g * 100, 1)}% · Vazias com IA > 35 dias: "
                 f"{br(v('Vazias com IA há > 35 dias') or 0)}")
    dif = v("Diferença DairyPlan × comercializado")
    if isinstance(dif, (int, float)):
        t.append(f"DairyPlan × leite comercializado: {'+' if dif > 0 else ''}{br(dif * 100, 1)}%")
    t.append("Relatório completo no PDF anexo.")
    return "\n".join(t)


def post(caminho, corpo):
    url = BASE.format(inst=os.environ["ZAPI_INSTANCE"], tok=os.environ["ZAPI_TOKEN"]) + caminho
    req = urllib.request.Request(url, data=json.dumps(corpo).encode(), method="POST",
                                 headers={"Content-Type": "application/json"})
    if os.environ.get("ZAPI_CLIENT_TOKEN"):
        req.add_header("Client-Token", os.environ["ZAPI_CLIENT_TOKEN"])
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.status, r.read().decode()[:300]


if __name__ == "__main__":
    falta = [k for k in ("ZAPI_INSTANCE", "ZAPI_TOKEN", "WPP_DESTINO") if not os.environ.get(k)]
    if falta:
        raise SystemExit(f"Faltam variáveis de ambiente: {', '.join(falta)}")
    texto = resumo()
    doc = "data:application/pdf;base64," + base64.b64encode(open(PDF, "rb").read()).decode()
    for fone in [re.sub(r"\D", "", f) for f in os.environ["WPP_DESTINO"].split(",") if f.strip()]:
        print(fone[:4] + "…", "texto:", *post("/send-text", {"phone": fone, "message": texto}))
        print(fone[:4] + "…", "pdf:", *post("/send-document/pdf", {
            "phone": fone, "document": doc, "fileName": os.path.basename(PDF)}))
