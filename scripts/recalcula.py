"""Recalcula todas as fórmulas de uma planilha .xlsx com o LibreOffice (sem interface) e confere se sobrou erro.

A planilha do monitoramento é gerada só com fórmulas (openpyxl não calcula). Este script abre o arquivo no
LibreOffice, manda recalcular tudo, salva por cima e depois procura células com erro (#VALOR!, #DIV/0! etc.).

Uso:
    python3 scripts/recalcula.py monitoramento/Monitoramento_Diario_ARCAFLA.xlsx [tempo_limite_s]
Saída: JSON com status ("success" ou "errors"), total de fórmulas e erros encontrados.
Requer: LibreOffice (comando soffice) e openpyxl.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub Recalcula(caminho As String)
    Dim args(0) As New com.sun.star.beans.PropertyValue
    args(0).Name = "Hidden"
    args(0).Value = True
    doc = StarDesktop.loadComponentFromURL(ConvertToURL(caminho), "_blank", 0, args())
    doc.calculateAll()
    doc.store()
    doc.close(True)
End Sub
</script:module>
"""
ERROS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:")


def soffice():
    for nome in ("soffice", "libreoffice"):
        if shutil.which(nome):
            return shutil.which(nome)
    raise SystemExit("LibreOffice (soffice) não encontrado; instale com: apt-get install -y libreoffice-calc")


def recalcula(arquivo, limite=600):
    arquivo = os.path.abspath(arquivo)
    perfil = tempfile.mkdtemp(prefix="lo_perfil_")
    env_perfil = f"-env:UserInstallation=file://{perfil}"
    so = soffice()
    # 1ª execução só cria o perfil; depois grava a macro no perfil
    subprocess.run([so, env_perfil, "--headless", "--terminate_after_init"], capture_output=True, timeout=120)
    pasta = os.path.join(perfil, "user", "basic", "Standard")
    os.makedirs(pasta, exist_ok=True)
    with open(os.path.join(pasta, "Module1.xba"), "w", encoding="utf-8") as f:
        f.write(MACRO)
    xlb = os.path.join(pasta, "script.xlb")
    if not os.path.exists(xlb):
        with open(xlb, "w", encoding="utf-8") as f:
            f.write('<?xml version="1.0" encoding="UTF-8"?>\n<!DOCTYPE library:library PUBLIC '
                    '"-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">\n<library:library '
                    'xmlns:library="http://openoffice.org/2000/library" library:name="Standard" '
                    'library:readonly="false" library:passwordprotected="false">\n'
                    ' <library:element library:name="Module1"/>\n</library:library>\n')
    subprocess.run([so, env_perfil, "--headless", "--norestore",
                    f'macro:///Standard.Module1.Recalcula("{arquivo}")'], capture_output=True, timeout=limite)
    shutil.rmtree(perfil, ignore_errors=True)


def confere(arquivo):
    formulas = sum(1 for ws in load_workbook(arquivo).worksheets for row in ws.iter_rows()
                   for c in row if isinstance(c.value, str) and c.value.startswith("="))
    erros = {}
    for ws in load_workbook(arquivo, data_only=True).worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROS):
                    erros.setdefault(c.value, []).append(f"{ws.title}!{c.coordinate}")
    total = sum(len(v) for v in erros.values())
    return {"status": "success" if total == 0 else "errors", "total_errors": total, "total_formulas": formulas,
            "error_summary": {k: {"count": len(v), "exemplos": v[:5]} for k, v in erros.items()}}


if __name__ == "__main__":
    arq = sys.argv[1]
    recalcula(arq, int(sys.argv[2]) if len(sys.argv) > 2 else 600)
    print(json.dumps(confere(arq), indent=2, ensure_ascii=False))
