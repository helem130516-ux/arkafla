"""Lê a exportação por vaca do OnFarm (dados/cultura_onfarm/*.xlsx, a mais recente) e gera:
- dados/cultura_onfarm_vacas.json: um registro por cultura (vaca, data, tipo, quarto, grau, resultado, agentes);
- dados/cultura_onfarm_mensal.json: nº de agentes isolados por mês em resultados "Positivo"
  (mesma contagem do gráfico do OnFarm). Meses que a planilha não cobre por inteiro no início
  vêm do gráfico transcrito (dados/cultura_onfarm_2026.json).
"""
import calendar
import glob
import json
from datetime import date

import openpyxl

MAP = {"STREP._AGALACTIAE_/_DYSGALACTIAE": "Strep. agalactiae / dysgalactiae", "STAPH._AUREUS": "Staph. aureus",
       "STREP._UBERIS": "Strep. uberis", "OUTROS_GRAM-POS": "Outros Gram-pos", "E._COLI": "E. coli",
       "KLEBSIELLA_/_ENTEROBACTER": "Klebsiella / Enterobacter", "PROTOTHECA_/_LEVEDURA": "Prototheca / Levedura",
       "OUTROS_GRAM-NEG": "Outros Gram-neg", "ENTEROCOCCUS_SPP": "Enterococcus spp", "LACTOCOCCUS_SPP": "Lactococcus spp",
       "PSEUDOMONAS_SPP": "Pseudomonas spp", "SERRATIA_SPP": "Serratia spp", "STAPH_NãO_AUREUS": "Staph não aureus"}

arq = sorted(glob.glob("dados/cultura_onfarm/*.xlsx"))[-1]
ws = openpyxl.load_workbook(arq, data_only=True).active
H = [c.value for c in ws[1]]
regs = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(H, row))
    if r["BRINCO"] in (None, ""):
        continue
    d, m, y = r["DATA"].split("/")
    v = str(r["BRINCO"]).strip()
    regs.append(dict(vaca=int(v) if v.isdigit() else v, data=f"{y}-{m}-{d}", tipo=r["TIPO"], leitura=r["TIPO_LEITURA"],
                     quarto=(r["QUARTO"] or "").replace(",", " "), grau=r["GRAU"], resultado=r["RESULTADO"],
                     agentes=[MAP[k] for k in MAP if r.get(k) == "Sim"], n_diag=r["DIAGNOSTICO POR ANIMAL"],
                     cura=r["CURA_CLINICA"], status=r["STATUS_VALIDACAO"]))
regs.sort(key=lambda x: (x["data"], str(x["vaca"])))
json.dump(dict(arquivo=arq.split("/")[-1], registros=regs), open("dados/cultura_onfarm_vacas.json", "w"), ensure_ascii=False, indent=0)

# mensal
P = json.load(open("dados/cultura_onfarm_2026.json"))
ag = P["agentes"]
d0, d1 = regs[0]["data"], regs[-1]["data"]
mes_ini = d0[:7] if d0[8:] <= "05" else None  # primeiro mês só vale se a planilha começa no início do mês
meses = sorted(set(P["meses"]) | {r["data"][:7] for r in regs})
n, fonte = {}, {}
for m in meses:
    usa_planilha = m > d0[:7] or (m == d0[:7] and mes_ini)
    if usa_planilha or m not in P["n"]:
        c = [0] * len(ag)
        for r in regs:
            if r["data"][:7] == m and r["resultado"] == "Positivo":
                for a in r["agentes"]:
                    c[ag.index(a)] += 1
        n[m], fonte[m] = c, "planilha"
    else:
        n[m], fonte[m] = P["n"][m], "gráfico"
y, mm = map(int, d1[:7].split("-"))
parcial = [d1[:7]] if int(d1[8:]) < calendar.monthrange(y, mm)[1] - 2 else []
json.dump(dict(agentes=ag, meses=meses, n=n, fonte=fonte, parcial=parcial,
               descricao=f"Agentes isolados em resultados positivos, por mês ({meses[0][5:]}/{meses[0][:4]} a {meses[-1][5:]}/{meses[-1][:4]}). "
                         f"Até {P['meses'][0][5:]}/{P['meses'][0][:4]}–{d0[5:7]}/{d0[:4]}: gráfico do OnFarm; depois: planilha por vaca ({arq.split('/')[-1]}, até {d1[8:]}/{d1[5:7]})."),
          open("dados/cultura_onfarm_mensal.json", "w"), ensure_ascii=False, indent=1)
print(len(regs), "culturas OnFarm por vaca", d0, "a", d1, "; mensal:", {m: (sum(n[m]), fonte[m]) for m in meses}, "parcial", parcial)
