"""Lê os boletins de cultura do LabVet (Carambeí) em dados/cultura_labvet e grava dados/cultura_labvet.json.

Identificação no boletim: "<quarto(s)> G<grau> L<lote>", ex.: "PD G1 L5".
"""
import glob
import json
import re
import warnings

import openpyxl

warnings.filterwarnings("ignore")
NEG = {"-", "", None}


def _vaca(v):
    s = str(v).strip()
    return int(float(s)) if re.fullmatch(r"\d+(\.0)?", s) else s


def le(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    ctrl, entrada, liberado, obs = ws["H3"].value, ws["H7"].value, None, []
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 6).value == "Conferido e liberado em: ":
            liberado = ws.cell(r, 7).value
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip().startswith("Obs"):
            obs = [str(ws.cell(r, c).value) for c in range(2, 13) if ws.cell(r, c).value not in NEG]
    hdr = next(r for r in range(1, 40) if ws.cell(r, 2).value == "AMOSTRA")
    out = []
    r = hdr + 1
    while ws.cell(r, 2).value not in NEG:
        ident = str(ws.cell(r, 5).value or "").strip()
        q = re.findall(r"\b(AE|AD|PE|PD)\b", ident)
        g = re.search(r"\bG(\d)\b", ident)
        lt = re.search(r"\bL(\d+)\b", ident)
        ag = [str(ws.cell(r, c).value).strip() for c in (6, 7, 8) if str(ws.cell(r, c).value).strip() not in NEG]
        out.append(dict(amostra=ws.cell(r, 2).value, data=ws.cell(r, 3).value.date().isoformat(),
                        vaca=_vaca(ws.cell(r, 4).value), ident=ident, quartos=" ".join(q),
                        grau=int(g.group(1)) if g else None, lote=int(lt.group(1)) if lt else None,
                        agentes=[a for a in ag if a != "Não houve crescimento"],
                        sem_crescimento="Não houve crescimento" in ag))
        r += 1
    return dict(arquivo=path.split("/")[-1], controle_lab=ctrl, entrada=entrada.date().isoformat(),
                liberado=liberado.date().isoformat() if liberado else None, obs=obs, amostras=out)


if __name__ == "__main__":
    bs = [le(p) for p in sorted(glob.glob("dados/cultura_labvet/LabVet_*.xlsx"))]
    json.dump(bs, open("dados/cultura_labvet.json", "w"), ensure_ascii=False, indent=1)
    for b in bs:
        print(b["entrada"], b["controle_lab"], len(b["amostras"]), "liberado", b["liberado"], b["obs"])
