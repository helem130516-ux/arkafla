"""Lê os Relatórios 2.2 da APCBRH, empilha os controles e classifica as vacas.

Uso: python3 scripts/processa_controles.py [pasta]  (padrão: dados/apcbrh)
Gera saida/controles.json e saida/registros.csv e imprime a conferência.
"""
import csv
import glob
import json
import math
import os
import sys
from datetime import datetime

import openpyxl

LIM = 200
DESVIO = "10"
TETO = 9998


def dt(s):
    s = (s or "").strip() if isinstance(s, str) else s
    if not s:
        return None
    if isinstance(s, datetime):
        return s.date()
    return datetime.strptime(s, "%d/%m/%Y").date()


def le_relatorio(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Dados"]
    meta, hdr = {}, None
    for r in range(1, 10):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for i, v in enumerate(vals):
            if isinstance(v, str) and v.strip().endswith(":"):
                nxt = next((x for x in vals[i + 1:] if x not in (None, "")), None)
                meta[v.strip()] = nxt
    for r in range(1, 20):
        if ws.cell(r, 2).value == "VACA":
            hdr = r
            break
    rows = []
    r = hdr + 1
    while ws.cell(r, 2).value not in (None, ""):
        g = lambda col: ws.cell(r, col).value
        rows.append(dict(
            seq=int(g(1)), vaca=str(g(2)).strip(), lote=str(g(3)).strip(),
            parto=dt(g(4)), del_=int(g(5) or 0), idade=g(6), prod=float(g(7) or 0),
            ccs=float(g(9)) if g(9) not in (None, "") else None,
            pct=g(11), ccs_tq=g(13)))
        r += 1
    return dict(
        arquivo=os.path.basename(path), reb=int(meta["REB:"]),
        n=int(meta["No ANIMAIS:"]), data=dt(meta["DATA DO CONTROLE:"]),
        emissao=meta.get("DATA DA EMISSÃO:"), prod_total=meta["PRODUÇÃO TOTAL:"],
        media_prod=meta["MEDIA PRODUÇÃO:"], ccs_tq=meta["MÉDIA CCS TANQUE:"], rows=rows)


def classifica(ctrls):
    ctrls.sort(key=lambda c: c["data"])
    recs = []
    for ci, c in enumerate(ctrls, 1):
        c["ci"] = ci
        for x in c["rows"]:
            recs.append(dict(x, ci=ci, data=c["data"]))
    recs.sort(key=lambda x: (int(x["vaca"]) if x["vaca"].isdigit() else 10**9, x["vaca"], x["ci"]))
    ult = {}  # vaca -> último registro (com estado da lactação)
    for x in recs:
        a = ult.get(x["vaca"])
        el = x["ccs"] is not None and x["ccs"] >= LIM
        x["prev_ccs"] = a["ccs"] if a else None
        nova_lac = bool(a and x["parto"] and x["parto"] > a["data"])
        consec = bool(a and a["ci"] == x["ci"] - 1)
        mesma_lac = bool(a and not nova_lac)
        if a is None or x["ccs"] is None:
            x["cls"] = "H"
        elif nova_lac:
            x["cls"] = "PE" if el else "PB"
        elif not consec or a["ccs"] is None:
            x["cls"] = "H"
        else:
            ael = a["ccs"] >= LIM
            x["cls"] = {(0, 0): "S", (0, 1): "N", (1, 1): "C", (1, 0): "R"}[(ael, el)]
        # já teve CCS elevada na mesma lactação (antes deste registro)
        teve_el = a["teve_el"] if mesma_lac else False
        x["rec"] = int(x["cls"] == "N" and teve_el)
        x["teve_el"] = teve_el or el
        x["seq_el"] = (a["seq_el"] + 1 if (consec and mesma_lac) else 1) if el else 0
        # queda: compara com o registro anterior da mesma lactação (aceita falta no controle)
        x["queda"] = int(bool(el and mesma_lac and a["prod"] > 0
                              and x["prod"] <= 0.7 * a["prod"] and x["del_"] >= 60))
        x["parto_alt"] = int(bool(a and a["parto"] and x["parto"] and x["parto"] != a["parto"]
                                  and x["parto"] < a["data"]))
        ult[x["vaca"]] = x
    return recs


def resumo(ctrls, recs):
    out = []
    for c in ctrls:
        rs = [x for x in recs if x["ci"] == c["ci"]]
        tp = sum(x["prod"] for x in rs)
        w = sum(x["prod"] * x["ccs"] for x in rs) / tp
        s10 = [x for x in rs if x["lote"] != DESVIO and x["ccs"] is not None]
        tp10 = sum(x["prod"] for x in s10)
        cnt = {k: sum(1 for x in rs if x["cls"] == k) for k in ["S", "N", "C", "R", "PB", "PE", "H"]}
        base = cnt["S"] + cnt["N"] + cnt["C"] + cnt["R"]
        ant = ctrls[c["ci"] - 2] if c["ci"] > 1 else None
        dias = (c["data"] - ant["data"]).days if ant else None
        tni = cnt["N"] / (cnt["S"] + cnt["N"]) if (cnt["S"] + cnt["N"]) else None
        vacas = {x["vaca"] for x in rs}
        vacas_ant = {x["vaca"] for x in recs if ant and x["ci"] == ant["ci"]}
        out.append(dict(
            ci=c["ci"], data=c["data"].isoformat(), reb=c["reb"], n_rel=c["n"], n_lin=len(rs),
            ccs_tq_apcbrh=round(c["ccs_tq"], 1), ccs_tq_calc=round(w, 1),
            ccs_tq_sem10=round(sum(x["prod"] * x["ccs"] for x in s10) / tp10, 1),
            media_arit_sem10=round(sum(x["ccs"] for x in s10) / len(s10), 1),
            media_geo_sem10=round(math.exp(sum(math.log(max(x["ccs"], 1)) for x in s10) / len(s10)), 1),
            n_lote10=sum(1 for x in rs if x["lote"] == DESVIO),
            pct_el=round(100 * sum(1 for x in rs if x["ccs"] >= LIM) / len(rs), 1),
            n_el=sum(1 for x in rs if x["ccs"] >= LIM),
            faixas=[sum(1 for x in rs if lo <= x["ccs"] < hi) for lo, hi in
                    [(0, 200), (200, 500), (500, 1000), (1000, 1e9)]],
            dias=dias, **cnt, base=base,
            pS=_p(cnt["S"], base), pN=_p(cnt["N"], base), pC=_p(cnt["C"], base), pR=_p(cnt["R"], base),
            taxa_ni=_p(cnt["N"], cnt["S"] + cnt["N"]),
            taxa_ni_30d=round(100 * tni * 30 / dias, 1) if (tni is not None and dias and dias >= 28) else None,
            permanencia=_p(cnt["C"], cnt["C"] + cnt["R"]),
            recorrentes=sum(x["rec"] for x in rs), recorrencia=_p(sum(x["rec"] for x in rs), cnt["N"]),
            cura_seco=_p(sum(1 for x in rs if x["cls"] in ("PB",) and x["prev_ccs"] is not None and x["prev_ccs"] >= LIM),
                         sum(1 for x in rs if x["cls"] in ("PB", "PE") and x["prev_ccs"] is not None and x["prev_ccs"] >= LIM)),
            ni_seco=_p(sum(1 for x in rs if x["cls"] == "PE" and x["prev_ccs"] is not None and x["prev_ccs"] < LIM),
                       sum(1 for x in rs if x["cls"] in ("PB", "PE") and x["prev_ccs"] is not None and x["prev_ccs"] < LIM)),
            seq3=sum(1 for x in rs if x["seq_el"] >= 3), queda=sum(x["queda"] for x in rs),
            del0=sum(1 for x in rs if x["del_"] == 0 or x["parto"] is None),
            teto=sum(1 for x in rs if x["ccs"] >= TETO),
            parto_alt=[x["vaca"] for x in rs if x["parto_alt"]],
            entraram=len(vacas - vacas_ant) if ant else None,
            sairam=len(vacas_ant - vacas) if ant else None,
            lotes=sorted({x["lote"] for x in rs}, key=lambda s: int(s) if s.isdigit() else 999),
        ))
    return out


def _p(a, b):
    return round(100 * a / b, 1) if b else None


if __name__ == "__main__":
    pasta = sys.argv[1] if len(sys.argv) > 1 else "dados/apcbrh"
    ctrls = [le_relatorio(p) for p in sorted(glob.glob(os.path.join(pasta, "R22_*.xlsx")))]
    recs = classifica(ctrls)
    res = resumo(ctrls, recs)
    os.makedirs("saida", exist_ok=True)
    with open("saida/controles.json", "w") as f:
        json.dump(res, f, ensure_ascii=False, indent=1)
    with open("saida/registros.csv", "w", newline="") as f:
        cols = ["ci", "data", "vaca", "lote", "parto", "del_", "prod", "ccs", "cls", "prev_ccs",
                "seq_el", "rec", "queda", "parto_alt", "ccs_tq"]
        wr = csv.DictWriter(f, cols, extrasaction="ignore")
        wr.writeheader()
        wr.writerows(recs)
    for r in res:
        print(json.dumps(r, ensure_ascii=False))
