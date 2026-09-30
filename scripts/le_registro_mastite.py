"""Lê dados/MASTITE_2026.xlsx (registro da fazenda) e grava dados/registro_fazenda.json:
tratamentos 2026 (com RETORNO), hiperqueratose (01/05/2026), secagens e doenças gerais."""
import datetime as dt
import json
import warnings

import openpyxl

warnings.filterwarnings("ignore")
wb = openpyxl.load_workbook("dados/MASTITE_2026.xlsx", data_only=True)
NUL = (None, "", "\xa0")
d = lambda x: x.date().isoformat() if isinstance(x, dt.datetime) else None  # noqa: E731

ws = wb["2026"]
trat = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 17)]
    if all(x in NUL for x in v[:14]):
        continue
    ret, obs = v[9], None
    if ret == "297/26":
        ret, obs = dt.datetime(2026, 7, 29), "retorno digitado '297/26', lido como 29/07/2026"
    q = " ".join(str(x).strip().upper() for x in v[10:14] if x not in NUL)
    trat.append(dict(linha=r, vaca=int(v[0]), ini=d(v[1]), produto=str(v[3]).strip(), dias=v[5], fim=d(v[7]),
                     carencia=v[8], retorno=d(ret), quartos=q, grau_hq=v[14], obs=obs))

ws = wb["HIPERQUERATOSE 26"]
hq_data = d(ws.cell(1, 1).value)
hq = {}
for r in range(3, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 8)]
    if isinstance(v[0], (int, float)) and any(isinstance(x, (int, float)) for x in v[1:5]):
        hq[int(v[0])] = dict(AD=v[1], AE=v[2], PE=v[3], PD=v[4], sujidade=v[5], media=v[6])

ws = wb["SECAGEM"]
sec = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 13)]
    if isinstance(v[1], (int, float)) and isinstance(v[5], dt.datetime):
        sec.append(dict(vaca=int(v[1]), lote=v[0], del_=v[2], lact=v[3], produto=str(v[4]).strip(), data=d(v[5]),
                        status=v[6], parto=d(v[8])))

ws = wb["DOENÇAS GERAL"]
doe = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 11)]
    if isinstance(v[0], dt.datetime) and v[1] not in NUL:
        doe.append(dict(data=d(v[0]), vaca=v[1], doenca=str(v[2]).strip().upper(), medicamento=str(v[3] or "").strip(),
                        retorno=d(v[9])))

json.dump(dict(tratamentos_2026=trat, hiperqueratose=dict(data=hq_data, vacas=hq), secagem=sec, doencas=doe),
          open("dados/registro_fazenda.json", "w"), ensure_ascii=False, indent=0)
print(len(trat), "tratamentos;", len(hq), "vacas com escore de hiperqueratose em", hq_data, ";", len(sec), "secagens;", len(doe), "doenças")
